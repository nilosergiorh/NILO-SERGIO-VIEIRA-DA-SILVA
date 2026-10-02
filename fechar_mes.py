#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera o arquivo TXT de importação do Domínio a partir da planilha do mês (já conferida).

Uso:   python fechar_mes.py PLANILHA.xlsx [--forcar]

1. Recalcula a planilha em segundo plano com o LibreOffice (se instalado); sem LibreOffice, usa os
   valores gravados pelo Excel na última vez que a planilha foi SALVA (salve antes de rodar).
2. Confere o semáforo de exportação (EXPORT_CALC!V11) e o leiaute das linhas.
3. Grava  lancamentos_dominio_AAAAMM.txt  (ANSI, uma linha por lançamento) na pasta da planilha.
Substitui o antigo "copiar EXPORT_TXT e colar no Bloco de Notas".
Sem --forcar, não gera o TXT enquanto houver pendências.
"""
import glob, os, shutil, subprocess, sys, tempfile, warnings

PASTA = os.path.dirname(os.path.abspath(__file__))


def soffice_python():
    """python.exe do LibreOffice (tem o módulo uno), se instalado."""
    for base in (os.environ.get('ProgramFiles', r'C:\Program Files'), os.environ.get('ProgramFiles(x86)', r'C:\Program Files (x86)'),
                 os.path.join(PASTA, 'LibreOfficePortable', 'App')):
        for p in glob.glob(os.path.join(base, 'LibreOffice*', 'program', 'python.exe')):
            return p
    return None


def recalcular(caminho, destino):
    """Recalcula todas as fórmulas no LibreOffice e salva uma cópia com os valores. Devolve True se conseguiu."""
    py = soffice_python()
    if not py: return False
    r = subprocess.run([py, os.path.join(PASTA, 'recalc_lo.py'), os.path.abspath(caminho), os.path.abspath(destino)],
                       capture_output=True, text=True, timeout=600)
    if r.returncode != 0 or not os.path.exists(destino):
        raise RuntimeError('LibreOffice não conseguiu recalcular a planilha:\n' + (r.stderr or r.stdout)[-800:])
    return True


def ler_exportacao(caminho):
    """Devolve (linhas_txt, semaforo, pendencias, competencia, origem_dos_valores)."""
    import openpyxl
    tmp = os.path.join(tempfile.mkdtemp(prefix='ponto_'), 'recalc.xlsx')
    try:
        origem = 'recalculado pelo LibreOffice' if recalcular(caminho, tmp) else 'valores salvos pelo Excel'
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            wb = openpyxl.load_workbook(tmp if os.path.exists(tmp) else caminho, data_only=True, read_only=False)
        ec = wb['EXPORT_CALC']
        semaforo = ec['V11'].value
        if semaforo is None:
            raise RuntimeError('A planilha não tem valores calculados. Abra-a no Excel, salve (Ctrl+S), feche e rode de novo '
                               '(ou instale o LibreOffice para o recálculo ser automático).')
        linhas = [str(c.value).strip() for c in wb['EXPORT_TXT']['A'] if c.value not in (None, '')]
        pend = {ec.cell(r, 21).value: ec.cell(r, 22).value for r in range(5, 11)}
        sheet, cell = wb.defined_names['P_COMP'].attr_text.replace('$', '').split('!')
        comp = str(wb[sheet.strip("'")][cell].value).split('.')[0]
        return linhas, str(semaforo), pend, comp, origem
    finally:
        shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)


def fechar(caminho, forcar=False):
    """Gera o TXT. Devolve (ok, mensagem, caminho_txt)."""
    linhas, semaforo, pend, comp, origem = ler_exportacao(caminho)
    resumo = '\n'.join(f'  - {k}: {v}' for k, v in pend.items() if k)
    if any('ERRO' in l for l in linhas):
        return False, ('Há códigos não preenchidos (empresa, rubrica ou empregado). '
                       'Complete PARAMETROS/FUNCIONARIOS e rode de novo.\n' + resumo), None
    if not semaforo.upper().startswith('PRONTO') and not forcar:
        return False, f'Semáforo: {semaforo}\n{resumo}\n\nResolva as pendências (aba PAINEL) e rode de novo.', None
    if not linhas:
        return False, 'Nenhum lançamento para exportar.', None
    ruins = [l for l in linhas if len(l) != 43 or not l.isdigit()]
    if ruins:
        return False, 'Linhas fora do leiaute (43 posições numéricas):\n' + '\n'.join(ruins[:5]), None
    destino = os.path.join(os.path.dirname(os.path.abspath(caminho)), f'lancamentos_dominio_{comp}.txt')
    with open(destino, 'w', encoding='cp1252', newline='\r\n') as f:
        f.write('\n'.join(linhas) + '\n')
    aviso = '' if semaforo.upper().startswith('PRONTO') else f'\nATENÇÃO: gerado com pendências ({semaforo}).'
    return True, f'{len(linhas)} lançamentos ({origem}) gravados em:\n{destino}{aviso}', destino


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    ok, msg, _ = fechar(sys.argv[1], '--forcar' in sys.argv)
    print(msg); sys.exit(0 if ok else 2)
