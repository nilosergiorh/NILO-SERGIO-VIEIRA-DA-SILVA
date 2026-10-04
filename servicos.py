# -*- coding: utf-8 -*-
"""Operações do sistema (usadas pelo aplicativo app.py e pelos atalhos ponto.py)."""
import datetime, json, os, re, shutil, subprocess, sys, tempfile, warnings

PASTA = os.path.dirname(os.path.abspath(__file__))
CLIENTES = os.path.join(PASTA, 'clientes')
sys.path.insert(0, PASTA)
import leitores, fechar_mes

R_MES = re.compile(r'^\d{4}-\d{2}$')


def D(s): return datetime.date.fromisoformat(s) if s else None


def clientes():
    return sorted(d for d in os.listdir(CLIENTES) if os.path.isdir(os.path.join(CLIENTES, d)) and not d.startswith('_'))


def config(cliente):
    return json.load(open(os.path.join(CLIENTES, cliente, 'config.json'), encoding='utf-8'))


def meses(cliente):
    base = os.path.join(CLIENTES, cliente)
    return sorted((d for d in os.listdir(base) if R_MES.match(d) and os.path.isdir(os.path.join(base, d))), reverse=True)


def planilha_do_mes(cliente, mes):
    pasta = os.path.join(CLIENTES, cliente, mes)
    pl = [os.path.join(pasta, f) for f in os.listdir(pasta) if f.startswith('Folha_Ponto_') and f.endswith('.xlsx')]
    return max(pl, key=os.path.getmtime) if pl else None


def livre(caminho):
    """Não sobrescreve: acrescenta _v2, _v3..."""
    base, ext = os.path.splitext(caminho); n = 2
    while os.path.exists(caminho):
        caminho = f'{base}_v{n}{ext}'; n += 1
    return caminho


def _do_mes(f):
    """Arquivos que pertencem à planilha atual do mês (planilha, cache do painel e TXT gerado dela)."""
    return (f.startswith('Folha_Ponto_') and (f.endswith('.xlsx') or f.endswith('.painel.json'))) or \
           (f.startswith('lancamentos_dominio_') and f.endswith('.txt'))


def substituir_planilha(pasta, nome_final, nova_temp, motivo):
    """Coloca `nova_temp` como a ÚNICA planilha do mês (`nome_final`). A planilha, o cache e o TXT anteriores
    vão para <mês>/_anteriores/<data hora>_<motivo>/ (nada é apagado). Devolve o caminho final."""
    antigos = [f for f in os.listdir(pasta) if _do_mes(f) and os.path.join(pasta, f) != nova_temp]
    for f in antigos:  # confere antes de mexer: planilha aberta no Excel não pode ser movida
        if f.endswith('.xlsx'):
            try:
                with open(os.path.join(pasta, f), 'r+b'): pass
            except PermissionError:
                os.remove(nova_temp)
                raise RuntimeError(f'A planilha "{f}" está aberta no Excel. Feche o Excel e importe de novo.')
    if antigos:
        dst = os.path.join(pasta, '_anteriores', datetime.datetime.now().strftime('%Y-%m-%d_%H%M%S') + '_' + motivo)
        os.makedirs(dst, exist_ok=True)
        for f in antigos: shutil.move(os.path.join(pasta, f), os.path.join(dst, f))
    final = os.path.join(pasta, nome_final)
    os.replace(nova_temp, final)
    return final


# ------------------------------------------------------------------ gerar a planilha do mês
def gerar_mes(cliente, arqs, log=print):
    """Lê os arquivos, gera a planilha do mês na pasta do cliente e devolve (caminho, mes, saida_texto)."""
    cfg = config(cliente)
    log(f'Lendo {len(arqs)} arquivo(s)...')
    recs = leitores.carregar(list(arqs), D(cfg.get('dt_ini')), D(cfg.get('dt_fim')), log=log)
    if not recs: raise RuntimeError('Nenhum registro de ponto encontrado nos arquivos.')
    ativos = [r['date'] for r in recs if r['punches']] or [r['date'] for r in recs]
    fim = D(cfg.get('dt_fim')) or max(ativos)
    comp = cfg.get('comp') or f'{fim.year}{fim.month:02d}'
    mes = f'{comp[:4]}-{comp[4:]}'
    pasta = os.path.join(CLIENTES, cliente, mes)
    os.makedirs(os.path.join(pasta, 'recebidos'), exist_ok=True)
    for a in arqs:  # arquiva o que o cliente mandou
        if os.path.dirname(os.path.abspath(a)).startswith(os.path.abspath(pasta)): continue
        dst = livre(os.path.join(pasta, 'recebidos', os.path.basename(a)))
        shutil.copy2(a, dst)
        if os.path.exists(a + '.leitura.json'): shutil.copy2(a + '.leitura.json', dst + '.leitura.json')
    nome = f'Folha_Ponto_{cliente.title()}_{comp[4:]}-{comp[:4]}.xlsx'
    temp = os.path.join(pasta, '_gerando_' + nome)  # só vira a planilha do mês se a geração der certo
    log('Montando a planilha...')
    r = subprocess.run([sys.executable, '-W', 'ignore', os.path.join(PASTA, 'gerar_folha.py'), ';'.join(arqs),
                        os.path.join(CLIENTES, cliente, 'config.json'), temp], capture_output=True, text=True, encoding='utf-8',
                       creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    if r.returncode != 0 or not os.path.exists(temp):
        if os.path.exists(temp): os.remove(temp)
        raise RuntimeError('Erro ao gerar a planilha:\n' + (r.stderr or r.stdout)[-1500:])
    saida = substituir_planilha(pasta, nome, temp, 'reimportacao')
    return saida, mes, r.stdout.replace(temp, saida)


# ------------------------------------------------------------------ dados do painel
def minutos(v):
    if v in (None, ''): return 0
    if isinstance(v, datetime.timedelta): return round(v.total_seconds() / 60)
    if isinstance(v, datetime.datetime): return round((v - datetime.datetime(1899, 12, 30)).total_seconds() / 60)
    if isinstance(v, datetime.time): return v.hour * 60 + v.minute
    if isinstance(v, (int, float)): return round(v * 1440)
    return leitores.hhmm(v) or 0


def _num(v):
    try: return int(round(float(v or 0)))
    except (TypeError, ValueError): return 0


def _data(v):
    d = leitores.data(v)
    return d.isoformat() if d else None


def _nome(wb, n):
    sheet, cell = wb.defined_names[n].attr_text.replace('$', '').split('!')
    return wb[sheet.strip("'")][cell].value


def dados_painel(planilha, recalcular=False, log=print):
    """Recalcula a planilha (LibreOffice) e extrai os números do painel. Guarda em cache até a planilha mudar."""
    import openpyxl
    cache = planilha + '.painel.json'
    mt = os.path.getmtime(planilha)
    if not recalcular and os.path.exists(cache):
        d = json.load(open(cache, encoding='utf-8'))
        if d.get('mtime') == mt: return d
    tmpdir = tempfile.mkdtemp(prefix='ponto_'); tmp = os.path.join(tmpdir, 'r.xlsx')
    try:
        log('Calculando (LibreOffice)...')
        fonte = tmp if fechar_mes.recalcular(planilha, tmp) else planilha
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            wb = openpyxl.load_workbook(fonte, data_only=True)
        ec = wb['EXPORT_CALC']
        if ec['V11'].value is None:
            raise RuntimeError('A planilha não tem valores calculados (instale o LibreOffice ou salve a planilha no Excel).')
        # por funcionário
        rs = wb['RESUMO_MENSAL']; funcs = []
        for r in range(5, rs.max_row + 1):
            nome = rs.cell(r, 1).value
            if not nome or nome == 'TOTAL': break
            c = lambda k: rs.cell(r, k).value
            funcs.append(dict(nome=nome, codigo=c(2) if c(2) not in ('', None) else None, status=str(c(25) or ''),
                              dias=_num(c(3)), impar=_num(c(4)), faltas_dias=_num(c(5)), trabalhado=minutos(c(7)),
                              previsto=minutos(c(8)), faltas_parc=minutos(c(17)), he70=minutos(c(19)), he100=minutos(c(20)),
                              noturno=minutos(c(21)), dsr=_num(c(16)), faltas_h=minutos(c(26)), dsr_h=minutos(c(27)), alertas=_num(c(24))))
        # pendências
        pd = wb['PENDENCIAS']; pend = []
        for r in range(6, pd.max_row + 1):
            if not pd.cell(r, 2).value: continue
            pend.append(dict(func=pd.cell(r, 2).value, data=_data(pd.cell(r, 3).value), dia=pd.cell(r, 4).value,
                             marc=pd.cell(r, 5).value or '', tipo=pd.cell(r, 6).value, confirmar=pd.cell(r, 7).value,
                             mensagem=pd.cell(r, 8).value))
        # irregularidades legais (CALCULO_DIARIO AM..AR) e alertas por dia
        ca = wb['CALCULO_DIARIO']
        tipos = [('jornada10', 39, 'Jornada acima de 10h (CLT art.59)'), ('interj', 40, 'Interjornada menor que 11h (CLT art.66)'),
                 ('almoco', 41, 'Almoço não registrado em jornada > 6h (CLT art.71)'), ('interv_menor', 42, 'Intervalo menor que 1h (CLT art.71)'),
                 ('interv_maior', 43, 'Intervalo maior que 2h (CLT art.71)'), ('ferias', 44, 'Trabalho em férias/afastamento')]
        irreg = {k: 0 for k, _, _ in tipos}; dias_irreg = []; ocorrencias = []
        for r in range(5, ca.max_row + 1):
            if not ca.cell(r, 1).value: continue
            for k, col, _ in tipos: irreg[k] += _num(ca.cell(r, col).value)
            # atestados, férias, licenças etc.: informação do dia (não são pendência)
            if ca.cell(r, 5).value and _num(ca.cell(r, 8).value):
                ocorrencias.append(dict(func=ca.cell(r, 1).value, data=_data(ca.cell(r, 2).value), dia=ca.cell(r, 3).value,
                                        ocorrencia=str(ca.cell(r, 5).value), grupo=str(ca.cell(r, 6).value or ''),
                                        trabalhado=minutos(ca.cell(r, 14).value), abonado=minutos(ca.cell(r, 19).value),
                                        alerta=ca.cell(r, 32).value or ''))
            if _num(ca.cell(r, 45).value):
                dias_irreg.append(dict(func=ca.cell(r, 1).value, data=_data(ca.cell(r, 2).value), alerta=ca.cell(r, 32).value or ''))
        pasta = os.path.dirname(planilha)
        comp = str(_nome(wb, 'P_COMP')).split('.')[0]
        txt = f'lancamentos_dominio_{comp}.txt'
        ativos = [f for f in funcs if not f['status'].startswith('IGNORADO')]
        d = dict(planilha=os.path.basename(planilha), mtime=mt, atualizado=datetime.datetime.now().isoformat(timespec='minutes'),
                 comp=comp, ini=_data(_nome(wb, 'P_DT_INI')), fim=_data(_nome(wb, 'P_DT_FIM')),
                 semaforo=str(ec['V11'].value), pronto=str(ec['V11'].value).upper().startswith('PRONTO'),
                 lancamentos=_num(ec['V5'].value), bloqueados=_num(ec['V9'].value), codigos_faltando=_num(ec['V10'].value),
                 totais=dict(funcionarios=len(ativos), he70=sum(f['he70'] for f in funcs), he100=sum(f['he100'] for f in funcs),
                             noturno=sum(f['noturno'] for f in funcs), faltas_h=sum(f['faltas_h'] for f in funcs),
                             faltas_dias=sum(f['faltas_dias'] for f in funcs), dsr=sum(f['dsr'] for f in funcs),
                             pendencias=len(pend), irregulares=len(dias_irreg)),
                 funcionarios=funcs, pendencias=pend, irreg=[dict(chave=k, nome=n, dias=irreg[k]) for k, _, n in tipos],
                 dias_irreg=dias_irreg[:300], ocorrencias=ocorrencias,
                 txt=txt if os.path.exists(os.path.join(pasta, txt)) else None)
        json.dump(d, open(cache, 'w', encoding='utf-8'), ensure_ascii=False)
        return d
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


# ------------------------------------------------------------------ mascote (assistente com IA)
MASCOTE = 'E-exato'


def _hm(m):
    m = round(m or 0); s = '-' if m < 0 else ''; m = abs(m)
    return f'{s}{m // 60}:{m % 60:02d}'


def contexto_mes(cliente, mes):
    """Resumo compacto do mês aberto, para o assistente responder com os números reais."""
    pl = planilha_do_mes(cliente, mes)
    if not pl or not os.path.exists(pl + '.painel.json'): return 'Nenhum mês aberto no painel.'
    d = json.load(open(pl + '.painel.json', encoding='utf-8'))
    cfg = config(cliente); t = d['totais']
    linhas = [f"Empresa: {cfg.get('empresa_nome', cliente)} (código Domínio {cfg.get('empresa')}). Competência {mes[5:]}/{mes[:4]}, "
              f"apuração {d['ini']} a {d['fim']}. Situação da exportação: {d['semaforo']}.",
              f"Regras: jornada {cfg.get('jornada')} seg-sex; tolerância {cfg.get('tolerancia')}/dia; HE a 70% até {cfg.get('he_lim')} no mês e 100% acima, "
              f"domingos e feriados; sábado a 100%: {cfg.get('sab100')}; falta de dia inteiro e DSR lançados como {cfg.get('horas_dia_falta')}; "
              f"rubricas {json.dumps(cfg.get('rubricas'), ensure_ascii=False)}.",
              f"Totais: HE70 {_hm(t['he70'])}, HE100 {_hm(t['he100'])}, adicional noturno {_hm(t['noturno'])}, faltas {_hm(t['faltas_h'])} "
              f"({t['faltas_dias']} dia(s) inteiro(s)), DSR descontado {t['dsr']}, pendências {t['pendencias']}, dias com irregularidade {t['irregulares']}.",
              'Por funcionário (nome | código | situação | HE70 | HE100 | noturno | faltas h | faltas dias | DSR | dias trabalhados):']
    cct = (cfg.get('cadastro') or {}).get('cct')
    if cct:  # cadastro da empresa: só consulta, não entra no cálculo
        linhas.insert(2, 'CCT cadastrada (informativo, o cálculo usa as regras acima): ' + '; '.join(f'{k}: {v}' for k, v in cct.items()) + '.')
    for f in d['funcionarios']:
        linhas.append(f"- {f['nome']} | {f['codigo']} | {f['status']} | {_hm(f['he70'])} | {_hm(f['he100'])} | {_hm(f['noturno'])} | "
                      f"{_hm(f['faltas_h'])} | {f['faltas_dias']} | {f['dsr']} | {f['dias']}")
    linhas.append('Pendências (confirmar com o cliente):')
    linhas += [f"- {p['mensagem']}" for p in d['pendencias']] or ['- nenhuma']
    linhas.append('Irregularidades legais (dias): ' + '; '.join(f"{i['nome']}: {i['dias']}" for i in d['irreg']))
    linhas += [f"- {x['func']} {x['data']}: {x['alerta']}" for x in d['dias_irreg'][:120]]
    linhas.append('Atestados e outras ocorrências do mês (informação, NÃO são pendência):')
    linhas += [f"- {o['func']} {o['data']}: {o['ocorrencia']} (trabalhou {_hm(o['trabalhado'])}, abonado {_hm(o['abonado'])})"
               for o in d.get('ocorrencias', [])] or ['- nenhuma']
    return '\n'.join(linhas)


SISTEMA = """Você é {nome}, o mascote-assistente do sistema de ponto da Exato Soluções Contábeis (escritório de contabilidade e departamento pessoal em Santa Catarina).
Você é um lobo-robô, versão digital do "Mestre Exato", cujo lema é "Mais que números, é estratégia para o seu futuro". Suas qualidades: inteligente (entende números e legislação), preciso (não trabalha com achismos), observador (encontra erros e oportunidades), estratégico (pensa antes de agir), confiável (sempre ao lado do cliente) e carismático (fala a língua das pessoas, com simplicidade e respeito). Fale de si no masculino; pode usar 🐺 com moderação.
Você conversa com a equipe do escritório, que entende de DP e da folha no Domínio, mas não de programação.
Responda em português do Brasil, curto e simpático, com linguagem de DP (sem termos de programação). Use listas curtas quando ajudar.
Use SOMENTE os números do mês abaixo quando falar de valores; se algo não estiver nos dados, diga que não tem essa informação.
Horas estão em horas:minutos. Quando pedirem mensagem para o cliente, escreva pronta para colar no WhatsApp, educada e objetiva.
Sobre legislação (CLT, CCT), responda com cautela e cite o artigo quando souber; lembre que a decisão final é do profissional.
Como o sistema funciona: o cliente manda o ponto (espelho, planilha, foto ou PDF) → "+ Novo mês" gera a planilha → pendências são confirmadas com o cliente e corrigidas no PONTO_BRUTO da planilha (Abrir planilha no Excel, salvar, Atualizar) → quando ficar "Pronto para importar", "Gerar TXT" → no Domínio: Utilitários > Importação > De Arquivo Texto > De Lançamentos.

DADOS DO MÊS ABERTO:
{contexto}"""


def conversar(cliente, mes, historico):
    """historico: [{'role': 'user'|'assistant', 'content': str}, ...] -> resposta do mascote."""
    chave = leitores.chave_api()
    if not chave: raise RuntimeError('sem_chave')
    import anthropic
    client = anthropic.Anthropic(api_key=chave)
    ctx = contexto_mes(cliente, mes) if cliente and mes else 'Nenhum mês aberto no painel.'
    msgs = [{'role': m['role'], 'content': str(m['content'])[:8000]} for m in historico[-20:] if m.get('role') in ('user', 'assistant')]
    while msgs and msgs[0]['role'] != 'user': msgs.pop(0)
    params = dict(model='claude-opus-5', max_tokens=16000, system=SISTEMA.format(nome=MASCOTE, contexto=ctx), messages=msgs)
    try:
        r = client.beta.messages.create(betas=['server-side-fallback-2026-07-01'], extra_body={'fallbacks': 'default'}, **params)
    except anthropic.BadRequestError:
        r = client.messages.create(**params)
    if r.stop_reason == 'refusal': return 'Desculpe, não consigo ajudar com isso.'
    return ''.join(b.text for b in r.content if b.type == 'text').strip()


def salvar_chave(k):
    k = (k or '').strip()
    if not k.startswith('sk-ant-'): raise ValueError('a chave da Anthropic começa com sk-ant-')
    open(os.path.join(PASTA, 'chave_api.txt'), 'w', encoding='utf-8').write(k)


def estado():
    """Clientes, meses e status resumido (do cache) para a tela inicial."""
    out = []
    for c in clientes():
        ms = []
        for m in meses(c):
            pl = planilha_do_mes(c, m)
            if not pl: continue
            cache = pl + '.painel.json'; st = None
            if os.path.exists(cache):
                d = json.load(open(cache, encoding='utf-8'))
                st = dict(pronto=d['pronto'], pendencias=d['totais']['pendencias'], txt=d.get('txt'),
                          desatualizado=d.get('mtime') != os.path.getmtime(pl))
            ms.append(dict(mes=m, planilha=os.path.basename(pl), status=st))
        out.append(dict(nome=c, empresa=config(c).get('empresa_nome', c), meses=ms))
    return out
