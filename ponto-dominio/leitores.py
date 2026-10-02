# -*- coding: utf-8 -*-
"""
Leitura dos arquivos de ponto enviados pelo cliente -> lista de registros (um por funcionário/dia).

Formatos aceitos:
  1. Espelho do relógio (Excel, uma aba por funcionário, colunas Ent./Saí. ... "Tipo de Cálculo").
  2. Planilha padrão: qualquer aba com cabeçalho Funcionário/Nome + Data + colunas de Entrada/Saída
     (inclui a aba PONTO_BRUTO de uma planilha já gerada - as correções e motivos são mantidos -
     e a planilha-modelo para clientes, gerada por  ponto.py modelo).
  3. Foto (JPG/PNG/WEBP), PDF ou planilha em outro formato: lidos pela IA (Claude), se houver chave
     da API em chave_api.txt ou na variável ANTHROPIC_API_KEY. O resultado da leitura fica salvo em
     <arquivo>.leitura.json ao lado do original (não paga duas vezes pela mesma leitura).

Registro: dict(emp, date, dow, punches[min], note, motivo, chprev, normais, faltas, atraso, extras,
               exsab, exdom, adnot, exfa, dsr)  - os campos do software ficam None fora do espelho.
"""
import base64, calendar, datetime, io, json, os, re, unicodedata

DOW = ['Seg', 'Ter', 'Qua', 'Qui', 'Sex', 'Sab', 'Dom']
SW = ('chprev', 'normais', 'faltas', 'atraso', 'extras', 'exsab', 'exdom', 'adnot', 'exfa', 'dsr')
MAX_MARC = 6
PASTA = os.path.dirname(os.path.abspath(__file__))


def norm(s):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFD', str(s or '')).encode('ascii', 'ignore').decode().upper()).strip()


def novo(emp, date, punches=(), note=None, motivo=None, dow=None, **sw):
    r = dict(emp=str(emp).strip(), date=date, dow=dow or DOW[date.weekday()], punches=list(punches),
             note=note or None, motivo=motivo or None)
    for k in SW: r[k] = sw.get(k)
    return r


def hhmm(v):
    """Horário/duração -> minutos (None se não for horário)."""
    if v is None or v == '': return None
    if isinstance(v, datetime.datetime):
        return v.hour * 60 + v.minute if v.year < 1901 else None
    if isinstance(v, datetime.time): return v.hour * 60 + v.minute
    if isinstance(v, datetime.timedelta): return round(v.total_seconds() / 60)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return round(v * 1440) if -1 < v < 2 else None
    if isinstance(v, str) and re.fullmatch(r'\d{1,3}:\d{2}(:\d{2})?', v.strip()):
        h, m = v.strip().split(':')[:2]; return int(h) * 60 + int(m)
    return None


def sgn(v):
    if v in (None, ''): return None
    if not isinstance(v, str): return hhmm(v)
    s = v.strip(); neg = s.startswith('-'); m = hhmm(s.lstrip('+-'))
    return None if m is None else (-m if neg else m)


def data(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    if isinstance(v, (int, float)) and 36526 <= v < 73051:
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
    m = re.search(r'(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})', str(v or ''))
    if m:
        y = int(m[3]); y += 2000 if y < 100 else 0
        try: return datetime.date(y, int(m[2]), int(m[1]))
        except ValueError: return None
    m = re.fullmatch(r'\s*(\d{4})-(\d{2})-(\d{2}).*', str(v or ''))
    return datetime.date(int(m[1]), int(m[2]), int(m[3])) if m else None


# ------------------------------------------------------------------ 1. espelho do relógio
def eh_espelho(wb):
    ws = wb.worksheets[0]
    cab = [str(ws.cell(1, c).value or '') for c in range(1, 32)]
    return 'Tipo de Cálculo' in cab or 'CHPrev' in cab


def ler_espelho(wb):
    recs = []
    for ws in wb:
        shifted = ws['D1'].value == 'Ent.2'
        hdr = {str(ws.cell(1, c).value): c for c in range(1, 32) if ws.cell(1, c).value}
        if 'CHPrev' not in hdr and 'Tipo de Cálculo' not in hdr: continue
        tc = hdr.get('Tipo de Cálculo', min((c for k, c in hdr.items() if k == 'CHPrev'), default=14))
        pcols = list(range(2, 14)) if shifted else [2, 3, 5, 6]
        merged = {}
        for mr in ws.merged_cells.ranges:
            v = ws.cell(mr.min_row, mr.min_col).value
            for r in range(mr.min_row, mr.max_row + 1): merged[r] = (v, mr.min_col)
        for r in range(3, ws.max_row + 1):
            a = ws.cell(r, 1).value
            d = data(a)
            if not d: continue
            punches = []; note = None
            for c in range(2, tc):
                v = ws.cell(r, c).value
                if v in (None, '') or (isinstance(v, str) and v.startswith('=')): continue
                if isinstance(v, datetime.time) and c not in pcols: continue  # auxiliar (08:45)
                m = hhmm(v) if not isinstance(v, (int, float)) else None
                if m is not None:
                    if c in pcols: punches.append(m)
                elif isinstance(v, str): note = v.strip()
            if r in merged and merged[r][0] and hhmm(merged[r][0]) is None: note = str(merged[r][0]).strip()
            g = lambda k: hhmm(ws.cell(r, hdr[k]).value) if k in hdr else None
            dow = a[11:14] if isinstance(a, str) and len(a) >= 14 else None
            rec = novo(ws.title, d, punches, note, dow=dow, chprev=g('CHPrev'), normais=g('Normais'), faltas=g('Faltas'),
                       atraso=g('Atraso'), extras=g('Extras'), exsab=g('ExSab'), exdom=g('ExDom'), adnot=g('Ad Not'),
                       exfa=sgn(ws.cell(r, hdr['Ex-Fa']).value) if 'Ex-Fa' in hdr else None, dsr=g('DSR'))
            if note and note.upper() == 'ATESTADO' and punches: rec['note'] = 'ATESTADO PARCIAL'
            recs.append(rec)
    return recs


# ------------------------------------------------------------------ 2. planilha padrão / PONTO_BRUTO
R_NOME = re.compile(r'^(funcion|nome|colaborador|empregado)', re.I)
R_DATA = re.compile(r'^data\b', re.I)
R_MARC = re.compile(r'^(ent(rada)?|sa[ií](da)?)\b\.?\s*\d*|^[es]\s?\d$', re.I)
R_OCOR = re.compile(r'observa|ocorr', re.I)
R_MOT = re.compile(r'tratamento|motivo', re.I)
R_SW = {'chprev': 'CHPrev', 'normais': 'Normais', 'faltas': 'Faltas', 'atraso': 'Atraso', 'extras': 'Extras',
        'exsab': 'ExSab', 'exdom': 'ExDom', 'adnot': 'Ad Not', 'exfa': 'Ex-Fa'}


def _cabecalho(ws):
    for r in range(1, 11):
        cab = [str(ws.cell(r, c).value or '').strip() for c in range(1, 41)]
        idx = lambda rx: next((i + 1 for i, t in enumerate(cab) if rx.search(t) and '(sw' not in t), None)
        marc = [i + 1 for i, t in enumerate(cab) if R_MARC.search(t) and '(sw' not in t]
        if idx(R_NOME) and idx(R_DATA) and len(marc) >= 2:
            sw = {k: next((i + 1 for i, t in enumerate(cab) if t.startswith(v + ' (sw')), None) for k, v in R_SW.items()}
            return r, dict(nome=idx(R_NOME), data=idx(R_DATA), marc=marc, ocor=idx(R_OCOR), mot=idx(R_MOT), sw=sw)
    return None, None


def ler_padrao(wb):
    recs = []
    abas = [wb['PONTO_BRUTO']] if 'PONTO_BRUTO' in wb.sheetnames else wb.worksheets
    for ws in abas:
        r0, col = _cabecalho(ws)
        if not r0: continue
        for r in range(r0 + 1, ws.max_row + 1):
            nome = ws.cell(r, col['nome']).value; d = data(ws.cell(r, col['data']).value)
            if not nome or not d: continue
            punches = [m for m in (hhmm(ws.cell(r, c).value) for c in col['marc']) if m is not None]
            note = ws.cell(r, col['ocor']).value if col['ocor'] else None
            mot = ws.cell(r, col['mot']).value if col['mot'] else None
            sw = {k: (sgn(ws.cell(r, c).value) if c else None) for k, c in col['sw'].items()}
            recs.append(novo(nome, d, punches, str(note).strip().upper() if note else None, str(mot).strip() if mot else None, **sw))
    return recs


# ------------------------------------------------------------------ 3. leitura com IA (Claude)
ESQUEMA = {'type': 'object', 'additionalProperties': False, 'required': ['funcionarios'], 'properties': {'funcionarios': {
    'type': 'array', 'items': {'type': 'object', 'additionalProperties': False, 'required': ['nome', 'dias'], 'properties': {
        'nome': {'type': 'string'}, 'dias': {'type': 'array', 'items': {
            'type': 'object', 'additionalProperties': False, 'required': ['data', 'marcacoes', 'ocorrencia', 'obs'], 'properties': {
                'data': {'type': 'string'}, 'marcacoes': {'type': 'array', 'items': {'type': 'string'}},
                'ocorrencia': {'type': 'string'}, 'obs': {'type': 'string'}}}}}}}}}

PROMPT = """Este arquivo é um registro/cartão/espelho de ponto de funcionários de uma empresa brasileira.
Extraia TODOS os dias de TODOS os funcionários presentes.{dica}
Para cada dia:
- data: AAAA-MM-DD (se o documento mostrar só o dia ou dia/mês, complete com o mês/ano do cabeçalho do documento).
- marcacoes: os horários batidos (entradas e saídas, incluindo saída e volta do almoço), em ordem cronológica, no formato HH:MM, exatamente como registrados. Não inclua colunas de totais, horas extras, faltas, saldo, banco de horas ou jornada prevista. Não invente horários: se algum estiver ilegível, omita-o e explique em obs.
- ocorrencia: texto da ocorrência escrito no documento para o dia (ex.: FALTA, FOLGA, FERIADO, ATESTADO, FÉRIAS, LICENÇA), em maiúsculas; "" se não houver.
- obs: observações relevantes (ex.: "horário rasurado", "assinatura ausente"), ou "".
Nome do funcionário exatamente como aparece no documento, em maiúsculas."""


def chave_api():
    k = os.environ.get('ANTHROPIC_API_KEY')
    arq = os.path.join(PASTA, 'chave_api.txt')
    if not k and os.path.exists(arq):
        k = open(arq, encoding='utf-8').read().strip()
    return k or None


def _bloco(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    raw = open(caminho, 'rb').read()
    if ext == '.pdf':
        return {'type': 'document', 'source': {'type': 'base64', 'media_type': 'application/pdf', 'data': base64.standard_b64encode(raw).decode()}}
    if ext in ('.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.tif', '.tiff', '.heic'):
        try:
            from PIL import Image, ImageOps
            img = ImageOps.exif_transpose(Image.open(io.BytesIO(raw))).convert('RGB')
            img.thumbnail((2400, 2400))
            buf = io.BytesIO(); img.save(buf, 'JPEG', quality=90); raw = buf.getvalue()
        except Exception as e:
            if ext not in ('.jpg', '.jpeg'): raise RuntimeError(f'não consegui abrir a imagem ({e}). Envie como JPG.')
        return {'type': 'image', 'source': {'type': 'base64', 'media_type': 'image/jpeg', 'data': base64.standard_b64encode(raw).decode()}}
    if ext in ('.xlsx', '.xlsm', '.xls', '.csv', '.txt'):
        if ext == '.csv' or ext == '.txt':
            txt = raw.decode('utf-8', errors='replace')
        else:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            partes = []
            for ws in wb:
                linhas = [';'.join('' if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
                partes.append(f'### Aba {ws.title}\n' + '\n'.join(l for l in linhas if l.strip(';')))
            txt = '\n\n'.join(partes)
        return {'type': 'text', 'text': txt}
    raise RuntimeError(f'tipo de arquivo não suportado: {ext}')


def ler_com_ia(caminho, dica=''):
    cache = caminho + '.leitura.json'
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(caminho):
        dados = json.load(open(cache, encoding='utf-8'))
    else:
        chave = chave_api()
        if not chave:
            raise RuntimeError('para ler foto/PDF é preciso a chave da API da Anthropic: salve-a em chave_api.txt na pasta do programa.')
        import anthropic
        client = anthropic.Anthropic(api_key=chave)
        params = dict(model='claude-opus-5', max_tokens=64000,
                      output_config={'format': {'type': 'json_schema', 'schema': ESQUEMA}},
                      messages=[{'role': 'user', 'content': [_bloco(caminho), {'type': 'text', 'text': PROMPT.format(dica=dica)}]}])
        try:
            with client.beta.messages.stream(betas=['server-side-fallback-2026-07-01'], extra_body={'fallbacks': 'default'}, **params) as s:
                msg = s.get_final_message()
        except anthropic.BadRequestError:
            with client.messages.stream(**params) as s:
                msg = s.get_final_message()
        if msg.stop_reason == 'refusal': raise RuntimeError('a IA recusou ler este arquivo.')
        if msg.stop_reason == 'max_tokens': raise RuntimeError('arquivo grande demais para uma leitura: divida em partes.')
        dados = json.loads(''.join(b.text for b in msg.content if b.type == 'text'))
        json.dump(dados, open(cache, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    recs = []
    for f in dados['funcionarios']:
        for d in f['dias']:
            dt = data(d['data'])
            if not dt: continue
            punches = [m for m in (hhmm(x) for x in d['marcacoes']) if m is not None]
            recs.append(novo(f['nome'].upper(), dt, punches, (d['ocorrencia'] or '').strip().upper() or None,
                             ('Lido por IA: ' + d['obs']) if d.get('obs') else None))
    return recs


# ------------------------------------------------------------------ carga geral
def ler_arquivo(caminho, dica=''):
    """Devolve (registros, formato)."""
    ext = os.path.splitext(caminho)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        import openpyxl, warnings
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            wb = openpyxl.load_workbook(caminho, data_only=False)
        if eh_espelho(wb): return ler_espelho(wb), 'espelho do relógio'
        recs = ler_padrao(wb)
        if recs: return recs, 'planilha padrão'
    return ler_com_ia(caminho, dica), 'leitura por IA'


def completar(recs, dt_ini=None, dt_fim=None):
    """Junta duplicados, preenche os dias sem registro e ordena por funcionário/data."""
    if not recs: return recs
    datas = [r['date'] for r in recs]
    ini = dt_ini or min(datas).replace(day=1)
    fim = dt_fim or max(datas).replace(day=calendar.monthrange(max(datas).year, max(datas).month)[1])
    ordem, por = [], {}
    for r in recs:
        if r['emp'] not in ordem: ordem.append(r['emp'])
        k = (r['emp'], r['date'])
        if k in por:
            a = por[k]
            a['punches'] += [m for m in r['punches'] if m not in a['punches']]
            a['note'] = a['note'] or r['note']; a['motivo'] = a['motivo'] or r['motivo']
            for s in SW: a[s] = a[s] if a[s] is not None else r[s]
        else:
            por[k] = dict(r)
    out = []
    for e in ordem:
        d = min(ini, min(r['date'] for r in recs if r['emp'] == e))
        ult = max(fim, max(r['date'] for r in recs if r['emp'] == e))
        while d <= ult:
            r = por.get((e, d)) or novo(e, d)
            if len(r['punches']) > MAX_MARC:
                extra = ' '.join(f'{m // 60:02d}:{m % 60:02d}' for m in r['punches'][MAX_MARC:])
                r['punches'] = r['punches'][:MAX_MARC]
                r['motivo'] = ((r['motivo'] or '') + f' Marcações além da 6ª ignoradas: {extra}.').strip()
            out.append(r); d += datetime.timedelta(days=1)
    return out


def carregar(caminhos, dt_ini=None, dt_fim=None, log=print):
    recs = []
    dica = ''
    if dt_ini and dt_fim:
        dica = f'\nPeríodo de apuração esperado: {dt_ini:%d/%m/%Y} a {dt_fim:%d/%m/%Y}.'
    for c in caminhos:
        rr, fmt = ler_arquivo(c, dica)
        log(f'  {os.path.basename(c)}: {fmt}, {len(rr)} dia(s), {len({r["emp"] for r in rr})} funcionário(s)')
        recs += rr
    return completar(recs, dt_ini, dt_fim)


def feriados_nacionais(ano):
    a = ano % 19; b, c = divmod(ano, 100); d, e = divmod(b, 4); f = (b + 8) // 25; g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30; i, k = divmod(c, 4); l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451; mes, dia = divmod(h + l - 7 * m + 114, 31)
    pascoa = datetime.date(ano, mes, dia + 1)
    fx = [(1, 1, 'Confraternização Universal'), (4, 21, 'Tiradentes'), (5, 1, 'Dia do Trabalho'),
          (9, 7, 'Independência do Brasil'), (10, 12, 'Nossa Senhora Aparecida'), (11, 2, 'Finados'),
          (11, 15, 'Proclamação da República'), (11, 20, 'Dia Nacional de Zumbi e da Consciência Negra'), (12, 25, 'Natal')]
    out = {datetime.date(ano, mm, dd): n for mm, dd, n in fx}
    out[pascoa - datetime.timedelta(days=2)] = 'Sexta-feira Santa'
    return out


def feriados(dt_ini, dt_fim, extras=None):
    """Feriados nacionais do período + extras do config ("AAAA-MM-DD" ou "DD/MM" anual; valor = descrição)."""
    out = {}
    for ano in range(dt_ini.year, dt_fim.year + 1):
        out.update(feriados_nacionais(ano))
        for k, v in (extras or {}).items():
            m = re.fullmatch(r'(\d{1,2})/(\d{1,2})', k.strip())
            d = datetime.date(ano, int(m[2]), int(m[1])) if m else data(k)
            if d: out[d] = v
    return {d: n for d, n in sorted(out.items()) if dt_ini <= d <= dt_fim}
