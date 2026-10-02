#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Atalhos do sistema de ponto -> Domínio Folha (chamado pelos arquivos .bat desta pasta).

  python ponto.py gerar    escolhe cliente + arquivos do mês -> gera a planilha e abre no Excel
  python ponto.py txt      escolhe a planilha conferida -> gera o TXT de importação do Domínio
  python ponto.py outros   planilha-modelo para o cliente, novo cliente, chave da IA
"""
import datetime, json, os, shutil, subprocess, sys, traceback
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog

PASTA = os.path.dirname(os.path.abspath(__file__))
CLIENTES = os.path.join(PASTA, 'clientes')
sys.path.insert(0, PASTA)
import leitores

root = tk.Tk(); root.withdraw(); root.attributes('-topmost', True)


def escolher(titulo, texto, opcoes):
    """Janela com lista de opções; devolve a escolhida ou None."""
    if not opcoes: return None
    win = tk.Toplevel(root); win.title(titulo); win.attributes('-topmost', True); win.resizable(False, False)
    tk.Label(win, text=texto, font=('Segoe UI', 10), padx=14, pady=10).pack(anchor='w')
    lb = tk.Listbox(win, height=min(12, len(opcoes)), width=48, font=('Segoe UI', 10), activestyle='dotbox')
    for o in opcoes: lb.insert('end', o)
    lb.selection_set(0); lb.pack(padx=14); lb.focus_set()
    res = {}
    def ok(*_):
        if lb.curselection(): res['v'] = opcoes[lb.curselection()[0]]
        win.destroy()
    lb.bind('<Double-Button-1>', ok); lb.bind('<Return>', ok)
    tk.Button(win, text='OK', width=12, command=ok).pack(pady=10)
    win.grab_set(); root.wait_window(win)
    return res.get('v')


def clientes():
    return sorted(d for d in os.listdir(CLIENTES) if os.path.isdir(os.path.join(CLIENTES, d)) and not d.startswith('_'))


def escolher_cliente():
    cs = clientes()
    return cs[0] if len(cs) == 1 else escolher('Cliente', 'Qual cliente?', cs)


def config(cliente):
    return json.load(open(os.path.join(CLIENTES, cliente, 'config.json'), encoding='utf-8'))


def D(s): return datetime.date.fromisoformat(s) if s else None


def livre(caminho):
    """Não sobrescreve: acrescenta _v2, _v3..."""
    base, ext = os.path.splitext(caminho); n = 2
    while os.path.exists(caminho):
        caminho = f'{base}_v{n}{ext}'; n += 1
    return caminho


# ------------------------------------------------------------------ 1. gerar planilha do mês
def gerar():
    cliente = escolher_cliente()
    if not cliente: return
    cfg = config(cliente)
    arqs = filedialog.askopenfilenames(
        parent=root, title=f'{cliente}: escolha os arquivos de ponto do mês (pode escolher vários)',
        initialdir=os.path.join(os.path.expanduser('~'), 'Downloads'),
        filetypes=[('Ponto (Excel, PDF, fotos)', '*.xlsx *.xlsm *.csv *.pdf *.jpg *.jpeg *.png *.webp *.heic'), ('Todos', '*.*')])
    if not arqs: return
    print(f'\nCliente: {cliente}')
    import servicos
    saida, mes, out = servicos.gerar_mes(cliente, list(arqs))
    print(out)
    avisos = [l.strip() for l in out.splitlines() if 'AVISO' in l or 'Feriados' in l]
    messagebox.showinfo('Ponto', f'Planilha gerada:\n{saida}\n\n' + '\n'.join(avisos) +
                        '\n\nEla vai abrir agora. Veja o PAINEL (pendências para confirmar com o cliente).')
    os.startfile(saida)


# ------------------------------------------------------------------ 2. gerar TXT
def txt():
    import fechar_mes
    arq = filedialog.askopenfilename(parent=root, title='Escolha a planilha do mês JÁ CONFERIDA (salva e fechada)',
                                     initialdir=CLIENTES, filetypes=[('Planilha do mês', 'Folha_Ponto_*.xlsx'), ('Excel', '*.xlsx')])
    if not arq: return
    print(f'\nRecalculando {os.path.basename(arq)}...')
    ok, msg, destino = fechar_mes.fechar(arq)
    if not ok and 'Semáforo' in msg:
        if messagebox.askyesno('Ponto', msg + '\n\nGerar o TXT MESMO ASSIM? (só se já conferiu as pendências)', icon='warning'):
            ok, msg, destino = fechar_mes.fechar(arq, forcar=True)
        else:
            return
    print(msg)
    if ok:
        messagebox.showinfo('Ponto', msg + '\n\nNo Domínio: Utilitários > Importação > De Arquivo Texto > De Lançamentos.')
        subprocess.run(['explorer', '/select,', os.path.normpath(destino)])
    else:
        messagebox.showerror('Ponto', msg)


# ------------------------------------------------------------------ 3. outros
def modelo():
    """Planilha-modelo para o cliente preencher (formato que o programa lê sem IA)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    cliente = escolher_cliente()
    if not cliente: return
    hoje = datetime.date.today(); padrao = f'{hoje.month:02d}/{hoje.year}'
    mes = simpledialog.askstring('Planilha-modelo', 'Mês de referência (MM/AAAA):', initialvalue=padrao, parent=root)
    if not mes: return
    m, a = (int(x) for x in mes.split('/'))
    cfg = config(cliente)
    nomes = []
    for d in sorted(os.listdir(os.path.join(CLIENTES, cliente)), reverse=True):  # funcionários da última planilha gerada
        pl = [f for f in os.listdir(os.path.join(CLIENTES, cliente, d)) if f.startswith('Folha_Ponto_')] if os.path.isdir(os.path.join(CLIENTES, cliente, d)) else []
        if pl:
            ws = openpyxl.load_workbook(os.path.join(CLIENTES, cliente, d, sorted(pl)[-1]), read_only=True)['FUNCIONARIOS']
            nomes = [r[0] for r in ws.iter_rows(min_row=5, max_col=7, values_only=True) if r[0] and r[6] != 'NÃO']
            break
    nomes = nomes or ['NOME DO FUNCIONÁRIO']
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'PONTO'
    ws['A1'] = f'Registro de ponto - {cfg.get("empresa_nome", cliente)} - {m:02d}/{a}'; ws['A1'].font = Font(size=14, bold=True)
    ws['A2'] = ('Preencha os horários no formato HH:MM (ex.: 08:00). Uma linha por funcionário/dia. '
                'Em faltas, atestados, férias, folgas etc., escreva a ocorrência na última coluna.')
    cab = ['Funcionário', 'Data', 'Dia', 'Entrada 1', 'Saída 1', 'Entrada 2', 'Saída 2', 'Entrada 3', 'Saída 3', 'Observação / ocorrência']
    for i, h in enumerate(cab, 1):
        c = ws.cell(4, i, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(horizontal='center')
    ultimo = (datetime.date(a + (m == 12), m % 12 + 1, 1) - datetime.timedelta(days=1)).day
    r = 5
    for n in nomes:
        for d in range(1, ultimo + 1):
            dt = datetime.date(a, m, d)
            ws.cell(r, 1, n); ws.cell(r, 2, dt).number_format = 'dd/mm/yyyy'; ws.cell(r, 3, leitores.DOW[dt.weekday()])
            for c in range(4, 10): ws.cell(r, c).number_format = 'hh:mm'
            if dt.weekday() >= 5:
                for c in range(1, 11): ws.cell(r, c).fill = PatternFill('solid', fgColor='EEEEEE')
            r += 1
    dv = DataValidation(type='list', formula1='"FALTA INJUSTIFICADA,FALTA JUSTIFICADA,ATESTADO,FÉRIAS,FOLGA COMPENSAÇÃO,FERIADO,AFASTAMENTO,ABONO,LICENÇA PATERNIDADE"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f'J5:J{r}')
    for col, w in zip('ABCDEFGHIJ', [34, 12, 6, 10, 10, 10, 10, 10, 10, 26]): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'D5'
    destino = livre(os.path.join(os.path.expanduser('~'), 'Downloads', f'Ponto_{cliente.title()}_{m:02d}-{a}_para_preencher.xlsx'))
    wb.save(destino)
    messagebox.showinfo('Ponto', f'Planilha-modelo salva em:\n{destino}\n\nEnvie ao cliente; quando voltar preenchida, use "1 Gerar planilha do mês".')
    subprocess.run(['explorer', '/select,', os.path.normpath(destino)])


def novo_cliente():
    nome = simpledialog.askstring('Novo cliente', 'Nome curto do cliente (ex.: EMPRESA_ABC):', parent=root)
    if not nome: return
    nome = leitores.norm(nome).replace(' ', '_')
    dst = os.path.join(CLIENTES, nome)
    if os.path.exists(dst):
        messagebox.showerror('Ponto', 'Esse cliente já existe.'); return
    os.makedirs(dst)
    cfg = json.load(open(os.path.join(CLIENTES, '_MODELO', 'config.json'), encoding='utf-8'))
    cfg['empresa_nome'] = nome.replace('_', ' ').title()
    json.dump(cfg, open(os.path.join(dst, 'config.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    messagebox.showinfo('Ponto', f'Cliente criado. Vou abrir o config.json no Bloco de Notas:\npreencha código da empresa, rubricas, jornada e códigos dos funcionários.')
    subprocess.Popen(['notepad', os.path.join(dst, 'config.json')])


def chave_ia():
    k = simpledialog.askstring('Chave da IA', 'Cole a chave da API da Anthropic (começa com sk-ant-).\n'
                               'Ela é usada só para ler fotos e PDFs e fica salva neste computador.', show='*', parent=root)
    if k:
        open(os.path.join(PASTA, 'chave_api.txt'), 'w', encoding='utf-8').write(k.strip())
        messagebox.showinfo('Ponto', 'Chave salva.')


def outros():
    op = escolher('Ponto', 'O que deseja fazer?', ['Planilha-modelo para o cliente preencher', 'Cadastrar novo cliente',
                                                   'Configurar chave da IA (fotos e PDFs)'])
    if op: {'Planilha': modelo, 'Cadastrar': novo_cliente, 'Configurar': chave_ia}[op.split()[0]]()


if __name__ == '__main__':
    acao = sys.argv[1] if len(sys.argv) > 1 else 'gerar'
    try:
        {'gerar': gerar, 'txt': txt, 'outros': outros}[acao]()
    except Exception as e:
        traceback.print_exc()
        messagebox.showerror('Ponto', f'Erro: {e}')
