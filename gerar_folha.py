#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador da planilha de apuração de ponto -> Folha Domínio.

Uso:   python gerar_folha.py ARQUIVO[;ARQUIVO2;...] CONFIG.json SAIDA.xlsx
       (arquivos: espelho do relógio, planilha padrão/PONTO_BRUTO já corrigido, foto ou PDF - veja leitores.py)
Requer: Python 3 e "pip install openpyxl" (+ "anthropic pillow" para foto/PDF).
Uso normal: dê dois cliques em "PONTO - 1 Gerar planilha do mes.bat".

O arquivo gerado contém fórmulas: abra no Excel (recalcula sozinho) ou envie ao Google Sheets.
O CONFIG.json guarda as regras do cliente (rubricas, códigos, jornada...) e as informações do mês
(férias, afastamentos, rescisões). Veja clientes/_MODELO/config.json.
"""
import sys, json, datetime, re, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import leitores

def D(s):
    return datetime.date.fromisoformat(s) if s else None
if len(sys.argv)<4:
    print(__doc__); sys.exit(1)
CFG=json.load(open(sys.argv[2],encoding='utf-8')); CFG['espelhos']=[p for p in sys.argv[1].split(';') if p.strip()]; CFG['saida']=sys.argv[3]

# ================= construção da planilha =================
import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

recs=leitores.carregar(CFG['espelhos'],D(CFG.get('dt_ini')),D(CFG.get('dt_fim')))
if not recs: sys.exit('Nenhum registro de ponto encontrado nos arquivos.')
def cfgm(d,nm):
    """Busca em um dict do config pelo nome completo ou pelo início do nome em palavras inteiras
    (ex.: 'MARIA' casa com 'MARIA DA SILVA SOUZA', mas 'RAFAEL' NÃO casa com 'RAFAELA ...').
    Ignora acentos; se várias chaves servirem, vale a mais longa."""
    if not d or not nm: return None
    N=leitores.norm(nm); best=None
    for k,v in d.items():
        K=leitores.norm(k)
        if K and (N==K or N.startswith(K+' ')) and (best is None or len(K)>best[0]): best=(len(K),v)
    return best[1] if best else None
CODES=CFG.get('codigos',{})
_dates=[x['date'] for x in recs]
DT_INI=D(CFG.get('dt_ini')) or min(_dates)
_act=[x['date'] for x in recs if x['punches']]   # última data com marcação real (ocorrências como FÉRIAS não contam)
DT_FIM=D(CFG.get('dt_fim')) or (max(_act) if _act else max(_dates))
COMP=CFG.get('comp') or f"{DT_FIM.year}{DT_FIM.month:02d}"
EMPNOME=CFG.get('empresa_nome','Empresa')
RUB=CFG.get('rubricas',{})
def hm(s,default):
    s=s or default; h_,m_=s.split(':'); return (int(h_)*60+int(m_))/1440
emps=[]
for r in recs:
    if r['emp'] not in emps: emps.append(r['emp'])
for _e in CFG.get('funcionarios_extra',[]):
    if _e not in emps: emps.append(_e)
NE=len(emps)
AVISOS=[]
for _k in CODES:
    _q=[e for e in emps if cfgm({_k:1},e) and cfgm(CODES,e)==CODES[_k]]
    if len(_q)>1: AVISOS.append(f'Código "{_k}" serve para {len(_q)} funcionários ({", ".join(_q)}): use o nome completo no config.')
for e in emps:
    if cfgm(CODES,e) is None and not cfgm(CFG.get('nao_controla'),e): AVISOS.append(f'{e}: sem código do Domínio no config.')
FERIADOS=leitores.feriados(DT_INI,DT_FIM,CFG.get('feriados'))
NFUNC=max(30,NE+10); FL=4+NFUNC   # linhas do cadastro FUNCIONARIOS (5..FL)
FIRST=5; LAST=FIRST+len(recs)-1        # linhas de dados PONTO_BRUTO / CALCULO
FONT='Arial'
f_norm=Font(name=FONT,size=10); f_bold=Font(name=FONT,size=10,bold=True)
f_hdr=Font(name=FONT,size=10,bold=True,color='FFFFFF')
f_in=Font(name=FONT,size=10,color='0000FF')
f_link=Font(name=FONT,size=10,color='008000')
f_title=Font(name=FONT,size=14,bold=True,color='1F3864')
f_note=Font(name=FONT,size=9,italic=True,color='595959')
fill_hdr=PatternFill('solid',fgColor='1F3864')
fill_crit=PatternFill('solid',fgColor='C55A11')
fill_in=PatternFill('solid',fgColor='FFFF00')
fill_sw=PatternFill('solid',fgColor='7F7F7F')
fill_sub=PatternFill('solid',fgColor='D9E1F2')
fill_warn=PatternFill('solid',fgColor='FCE4D6')
thin=Side(style='thin',color='BFBFBF'); box=Border(left=thin,right=thin,top=thin,bottom=thin)
T_HM='[h]:mm;-[h]:mm;'      # zero some (fica em branco)
T_HMZ='[h]:mm;-[h]:mm;0:00'
T_CLK='hh:mm'; T_DT='dd/mm/yyyy'

wb=Workbook()
def sheet(name,tab=None):
    ws=wb.create_sheet(name)
    if tab: ws.sheet_properties.tabColor=tab
    return ws
wb.remove(wb.active)
ws_lm=sheet('LEIA_ME','1F3864'); ws_dic=sheet('DICIONARIO','1F3864')
ws_par=sheet('PARAMETROS','FFC000'); ws_fun=sheet('FUNCIONARIOS','FFC000'); ws_tab=sheet('TABELAS','FFC000')
ws_pb=sheet('PONTO_BRUTO','7F7F7F'); ws_ca=sheet('CALCULO_DIARIO','2E75B6'); ws_se=sheet('SEMANAL','2E75B6')
ws_re=sheet('RESUMO_MENSAL','548235'); ws_cf=sheet('CONFERENCIA','C00000'); ws_dm=sheet('DEMONSTRACAO','C00000')
ws_ec=sheet('EXPORT_CALC','7030A0'); ws_csv=sheet('EXPORT_CSV','7030A0'); ws_txt=sheet('EXPORT_TXT','7030A0')

def hdr(ws,row,col,text,crit=False,sw=False,width=None):
    c=ws.cell(row,col,('★ ' if crit else '')+text)
    c.font=f_hdr; c.fill=fill_crit if crit else (fill_sw if sw else fill_hdr)
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=box
    if width: ws.column_dimensions[L(col)].width=width
def title(ws,text,sub=None):
    ws['A1']=text; ws['A1'].font=f_title
    if sub: ws['A2']=sub; ws['A2'].font=f_note
def name(n,ref):
    wb.defined_names[n]=DefinedName(n,attr_text=ref)

# ================= PARAMETROS =================
ws=ws_par
title(ws,'PARAMETROS - regras de negócio (única fonte das regras)','Células AMARELAS/azuis = editáveis. Todas as fórmulas do sistema leem daqui pelos nomes definidos (coluna A). Alterou aqui, recalcula tudo.')
for i,h in enumerate(['Nome (fórmulas)','Parâmetro','Valor','Unidade','Base / origem da regra','Como customizar / observação'],1): hdr(ws,4,i,h,width=[16,44,14,10,52,70][i-1])
T=lambda h,m=0: (h*60+m)/1440
params=[
 ('P_DT_INI','Início da apuração',DT_INI,'data',T_DT,'Primeira data do espelho enviado pelo cliente.','CCT cl.31 permite também fechar de 21→20 ou 16→15; se usar, altere as datas.'),
 ('P_DT_FIM','Fim da apuração (dias após esta data NÃO são calculados)',DT_FIM,'data',T_DT,'Última data com marcação no espelho (ou "dt_fim" do config, quando o RH informar outro fechamento).','Se o cliente fechar em outra data, ajuste aqui.'),
 ('P_COMP','Competência (AAAAMM)',COMP,'texto',None,'Competência da folha.','Formato AAAAMM (usado no TXT posicional).'),
 ('P_EMPRESA','Código da empresa no Domínio',CFG.get('empresa'),'nº','0',CFG.get('empresa_origem') or 'Código da empresa no Domínio ("empresa" no config).','Sem ele o TXT mostra erro (evita importar na empresa errada).'),
 ('P_TIPOFOLHA','Tipo da folha na importação',CFG.get('tipo_folha',11),'cód.',None,'Domínio: 11-Mensal; 41-Adiantamento; 42-Complementar; 51/52-13º.','Mantenha 11 para folha mensal.'),
 ('P_JORNADA','Jornada diária contratual (seg-sex)',hm(CFG.get('jornada'),'08:45'),'hh:mm',T_HM,'Informado pelo RH: 08:45 seg-sex (08:45 x 5 = 43:45, dentro das 44h). ATENÇÃO: o espelho (CHPrev) está configurado com 08:30, divergindo da jornada real. CCT cl.24: prorrogação seg-sex compensa o sábado.','Individual: FUNCIONARIOS col.C. Se a empresa NÃO tem compensação de sábado, use 08:00 (limite CLT art.58/59).'),
 ('P_LIM_SEM','Limite semanal legal',T(44,0),'hh:mm',T_HM,'CF art.7º XIII / CCT cl.3 §2 (44h semanais).','Só usado na verificação SEMANAL (alerta se semana > 44h).'),
 ('P_TOL','Tolerância diária (atraso/saída antecipada/extra)',hm(CFG.get('tolerancia'),'00:10'),'hh:mm',T_HM,'CLT art.58 §1º: variações até 5 min por marcação, máx. 10 min/dia, não são descontadas nem pagas.','Ponha 00:00 para apurar minuto a minuto (o software do espelho apura sem tolerância).'),
 ('P_LIM_DIA','Limite de trabalho diário (alerta)',T(10,0),'hh:mm',T_HM,'CLT art.59: máx. 2h extras/dia (8h+2h). Aqui 10:00 = alerta.','Apenas sinaliza; não altera cálculo.'),
 ('P_INTERV_LIM','Jornada que exige intervalo (dia com só 1 par de marcações)',T(6,0),'hh:mm',T_HM,'CLT art.71: acima de 6h exige intervalo.','Se o dia tem só 1 par e passa disso, aplica o intervalo pré-assinalado abaixo.'),
 ('P_INTERV_PRE','Intervalo de almoço padrão da empresa',T(1,0),'hh:mm',T_HM,'Regra da empresa (informada pelo RH): 01:00. Usado só quando P_INTERV_MODO = 1.','Com marcação completa (4 batidas) o cálculo usa as batidas reais: almoço menor que 01:00 já vira hora extra automaticamente.'),
 ('P_INTERV_MODO','Dia com só 1 par de marcações (sem almoço batido)',CFG.get('interv_modo',2),'1/2','0','Regra da empresa: se o intervalo não foi respeitado (trabalhou), o tempo é pago como hora extra. 2 = NÃO deduz o almoço (paga como extra). 1 = deduz P_INTERV_PRE (só se o almoço foi feito e faltou bater).','Padrão = 2 (regra informada). Se a conferência mostrar que foi só esquecimento de batida, use 1 - e o sistema deduz 01:00 desses dias.'),
 ('P_INTERV_MIN','Intervalo mínimo em jornada > 6h (alerta)',T(1,0),'hh:mm',T_HM,'CLT art.71: mínimo de 1h (salvo redução por norma coletiva/autorização do MTE - verificar CCT).','Apenas sinaliza no PAINEL.'),
 ('P_INTERV_MAX','Intervalo máximo (alerta)',T(2,0),'hh:mm',T_HM,'CLT art.71: máximo de 2h (salvo acordo escrito/norma coletiva).','Apenas sinaliza no PAINEL.'),
 ('P_INTERJ','Interjornada mínima (alerta)',T(11,0),'hh:mm',T_HM,'CLT art.66: 11h entre jornadas.','Apenas sinaliza.'),
 ('P_NOT_INI','Início do horário noturno',T(22,0),'hh:mm',T_CLK,'CCT cl.8 (22:00 às 05:00).','—'),
 ('P_NOT_FIM','Fim do horário noturno',T(5,0),'hh:mm',T_CLK,'CCT cl.8.','—'),
 ('P_HORA_NOT','Duração da hora noturna',52.5/1440,'hh:mm:ss','hh:mm:ss','CLT art.73 §1º: hora noturna = 52min30s (CCT cl.8 é omissa; aplica-se a CLT).','Se a empresa NÃO reduz a hora noturna, ponha 01:00:00.'),
 ('P_NOT_ADIC','Adicional noturno',0.35,'%','0%','CCT cl.8: 35% (a CLT prevê 20%; vale a CCT).','Percentual é configurado na rubrica do Domínio; aqui é documentação.'),
 ('P_HE1','Adicional HE - faixa 1',0.70,'%','0%','CCT cl.7 "a": até 30 horas = 70%.','Documentação (percentual vive na rubrica do Domínio).'),
 ('P_HE2','Adicional HE - faixa 2 / domingos e feriados',1.00,'%','0%','CCT cl.7 "b" e "c": acima de 30h = 100%; domingos e feriados = 100%.','Documentação.'),
 ('P_HE_LIM','Horas extras na faixa de 70% (por mês)',hm(CFG.get('he_lim'),'30:00'),'hh:mm',T_HM,'CCT cl.7 "a" + regra confirmada pelo RH: 30h por MÊS a 70%; acima disso, 100%. Conta extras de dia útil + sábado (domingo/feriado sempre 100%).','Ajuste se a regra mudar.'),
 ('P_DSR_MEIO','Falta em 1 dia que faz perder o DSR (meio período)','=P_JORNADA/2','hh:mm',T_HM,'Regra da empresa (informada pelo RH): a partir de meio período de falta no dia, desconta o DSR da semana. Padrão = metade da jornada (04:22). Falta integral sempre perde. Obs.: é regra mais favorável ao empregado que a Lei 605/49 art.6º e a CCT cl.27 (atrasos até 30 min).','Digite outro horário para mudar a regra (ex.: 00:30 para seguir a CCT cl.27). Só falta INJUSTIFICADA conta: atestados e abonos não contam.'),
 ('P_MODO','Modo faltas x extras (1 ou 2)',1,'1/2','0','1 = paga extras e desconta faltas separadamente (padrão sem banco de horas). 2 = compensa dentro do mês (só saldo). CCT cl.29: banco de horas exige acordo/consulta.','O espelho traz "Ex-Fa" (saldo), o que sugere compensação; só use 2 se houver acordo válido.'),
 ('P_FATOR','Fator do valor no TXT (valor x fator)',100,'nº','0','Manual Domínio (Importar Lançamentos): valor com 2 casas decimais implícitas (ex.: 33,33 -> 000003333). Rubricas de horas: horas decimais x 100 (30:00 -> 3000).','Na tela de importação, escolha a opção de conversão de horas (campo "Lançamento de horas") compatível com horas decimais e TESTE com 1 empregado antes de importar tudo.'),
 ('R_HE70','Cód. rubrica Domínio - HE 70%',RUB.get('HE70'),'cód.','0','Movimentos Domínio: 170 HORAS EXTRAS 70% (unidade Horas).','—'),
 ('R_HE100','Cód. rubrica Domínio - HE 100% (>30h e dom/fer)',RUB.get('HE100'),'cód.','0','Movimentos Domínio: 200 Hora Extra 100% (unidade Horas).','—'),
 ('R_NOT','Cód. rubrica Domínio - Adicional noturno (horas)',RUB.get('NOT'),'cód.','0','Movimentos Domínio: 26 ADICIONAL NOTURNO automático (unidade Horas). ATENÇÃO: no relatório o adicional sai a 20%, e a CCT cl.8 prevê 35% - conferir o percentual da rubrica.','—'),
 ('R_NOTRED','Cód. rubrica Domínio - Hora noturna reduzida',RUB.get('NOTRED'),'cód.',None,'Não há rubrica própria no relatório Movimentos; deixe em branco (a redução não é exportada).','Se o Domínio tiver rubrica de redução, informe o código.'),
 ('R_FALTA','Cód. rubrica Domínio - Faltas em horas (parciais + dias inteiros)',RUB.get('FALTA'),'cód.','0','Movimentos Domínio: 8069 HORAS FALTAS PARCIAL (unidade Horas, desconto).','Faltas de dia inteiro entram como horas (P_HORAS_DIA_FALTA por dia) nesta mesma rubrica - CONFIRMAR como o RH lança hoje.'),
 ('R_DSR','Cód. rubrica Domínio - DSR perdido (horas)',RUB.get('DSR'),'cód.','0','Movimentos Domínio: 42 HORAS FALTAS dsr (unidade Horas; nos meses anteriores lançado como 7:20 por DSR perdido).','—'),
 ('P_HORAS_DIA_FALTA','Horas de 1 dia (falta de dia inteiro e DSR)',hm(CFG.get('horas_dia_falta'),'07:20'),'hh:mm',T_HM,'Confirmado pelo RH: no espelho a jornada é 08:45, mas o Domínio considera 1 dia de trabalho = 07:20 (220:00/30). Falta de dia inteiro e DSR perdido são lançados em 07:20. Faltas parciais/atrasos continuam em horas reais.','Altere se a empresa mudar a carga horária do dia no Domínio.'),
 ('P_SAB100','Sábado trabalhado pago a 100%? (SIM/NÃO)',CFG.get('sab100','SIM'),'SIM/NÃO',None,'Confirmado pelo RH: sábado é dia de descanso (já compensado durante a semana), então o trabalho no sábado é hora extra a 100%. É mais favorável que a CCT cl.7 (que prevê 100% só para domingos e feriados). Coerente com os lançamentos de 08/2026 no Domínio.','SIM = extras de sábado vão direto para HE 100% e não contam na faixa de 30h a 70%. NÃO = seguem a faixa 70%/100% dos dias úteis.'),
]
prow={}
for i,(n,lab,val,un,fmt,base,how) in enumerate(params):
    r=5+i; prow[n]=r
    ws.cell(r,1,n).font=Font(name=FONT,size=9,bold=True)
    ws.cell(r,2,lab).font=f_norm
    c=ws.cell(r,3,val); c.font=f_in; c.fill=fill_in; c.border=box
    if fmt: c.number_format=fmt
    ws.cell(r,4,un).font=f_norm
    ws.cell(r,5,base).font=f_norm; ws.cell(r,6,how).font=f_norm
    for k in (2,5,6): ws.cell(r,k).alignment=Alignment(wrap_text=True,vertical='top')
    name(n,f"PARAMETROS!$C${r}")
ws.freeze_panes='A5'

# ================= TABELAS =================
ws=ws_tab
title(ws,'TABELAS - ocorrências, feriados','Cada texto escrito na coluna "Observação" do PONTO_BRUTO precisa existir aqui. Ocorrência nova? Acrescente uma linha (amarelo).')
for i,h in enumerate(['Ocorrência (texto no ponto)','Grupo','Abona falta? (SIM/NÃO)','Base legal / observação'],1): hdr(ws,4,i,h,width=[30,18,14,70][i-1])
oc=[('ATESTADO','ATESTADO','SIM','CLT art.473/Lei 605 art.6º §2º; CCT cl.35 (entregar em 24h úteis). Até 15 dias pago pela empresa.'),
    ('ATESTADO PARCIAL','ATESTADO','SIM','Atestado em horas: CCT cl.35 §2º acrescenta 1h antes e 1h depois como falta justificada. Abona só o saldo faltante do dia.'),
    ('FÉRIAS','FÉRIAS','SIM','Lançar no módulo Férias do Domínio (não via ponto). Sem desconto.'),
    ('LICENÇA PATERNIDADE','LICENÇA','SIM','CCT cl.28 (nascimento: 5 dias). Sem desconto.'),
    ('FALTA JUSTIFICADA','FALTA JUST.','SIM','CCT cl.28 (casamento, óbito, internação, vestibular, CNH...).'),
    ('AFASTAMENTO','AFASTAMENTO','SIM','Afastamento (INSS/outros) - lançar no Domínio.'),
    ('AFASTAMENTO INSS','AFASTAMENTO','SIM','Afastamento previdenciário - lançar no Domínio.'),
    ('ACIDENTE DE TRABALHO','AFASTAMENTO','SIM','Afastamento por acidente de trabalho - lançar no Domínio (CAT/benefício acidentário).'),
    ('FOLGA COMPENSAÇÃO','COMPENSAÇÃO','SIM','Compensação de feriado/ponte (CCT cl.25).'),
    ('ABONO','ABONO','SIM','Abono gerencial - documentar autorização.'),
    ('FALTA INJUSTIFICADA','FALTA','NÃO','Desconta o dia e o DSR da semana (Lei 605/49 art.6º).'),
    ('FERIADO','FERIADO','SIM','Feriado marcado no próprio espelho. Cadastre a data na lista de feriados (col. F) para o trabalho no dia virar HE 100%.')]
for i in range(26):
    r=5+i
    for j in range(4):
        c=ws.cell(r,j+1,oc[i][j] if i<len(oc) else None); c.font=f_in if j<3 else f_norm; c.border=box
        if i>=len(oc) or j<3: c.fill=fill_in
        if j==3: c.alignment=Alignment(wrap_text=True,vertical='top')
name('TB_OC_NOME','TABELAS!$A$5:$A$30'); name('TB_OC_GRUPO','TABELAS!$B$5:$B$30'); name('TB_OC_ABONA','TABELAS!$C$5:$C$30')
hdr(ws,4,6,'Feriados (datas)',width=16); hdr(ws,4,7,'Feriado',width=40)
ws.cell(3,6,'Feriados do período: nacionais preenchidos automaticamente + estaduais/municipais do config ("feriados"). Acrescente outros nas linhas amarelas vazias.').font=f_note
_fer=list(FERIADOS.items())
for i in range(20):
    c=ws.cell(5+i,6,_fer[i][0] if i<len(_fer) else None); c.number_format=T_DT; c.font=f_in; c.fill=fill_in; c.border=box
    c=ws.cell(5+i,7,_fer[i][1] if i<len(_fer) else None); c.font=f_norm; c.border=box
name('TB_FERIADOS','TABELAS!$F$5:$F$24')

# ================= FUNCIONARIOS =================
ws=ws_fun
title(ws,'FUNCIONARIOS - cadastro e códigos do Domínio','Nome deve ser IDÊNTICO ao usado no PONTO_BRUTO (nome da aba do espelho). Código Domínio = código do empregado na Folha.')
for i,h in enumerate(['Funcionário (como no espelho)','Código Domínio','Jornada diária (opcional)','Férias - início','Férias - fim','Data de rescisão (último dia trabalhado)','Controla ponto? (SIM/NÃO)','Observação','Afastamento - início','Afastamento - fim (vazio = em aberto)','Tipo do afastamento'],1): hdr(ws,4,i,h,width=[36,16,18,14,14,20,16,60,14,18,24][i-1])
ws.row_dimensions[4].height=44
for i in range(NFUNC):
    r=5+i
    nm=emps[i] if i<NE else None
    c=ws.cell(r,1,nm); c.font=f_in; c.fill=fill_in; c.border=box
    c=ws.cell(r,2,(cfgm(CODES,nm) if nm else None)); c.font=f_in; c.fill=fill_in; c.border=box
    c=ws.cell(r,3,'=P_JORNADA' if nm else None); c.number_format=T_HM; c.font=f_in; c.fill=fill_in; c.border=box
    for k in (4,5,6,9,10):
        c=ws.cell(r,k); c.number_format=T_DT; c.font=f_in; c.fill=fill_in; c.border=box
    c=ws.cell(r,11); c.font=f_in; c.fill=fill_in; c.border=box
    obs_=[]
    nc=cfgm(CFG.get('nao_controla'),nm) if nm else None
    c=ws.cell(r,7,('NÃO' if nc else 'SIM') if nm else None); c.font=f_in; c.fill=fill_in; c.border=box
    if nc: obs_.append(nc)
    fe=cfgm(CFG.get('ferias'),nm) if nm else None
    if fe:
        ws.cell(r,4,D(fe['inicio'])); ws.cell(r,5,D(fe['fim']))
        if fe.get('obs'): obs_.append(fe['obs'])
    af=cfgm(CFG.get('afastamentos'),nm) if nm else None
    if af:
        ws.cell(r,9,D(af['inicio'])); ws.cell(r,10,D(af.get('fim'))); ws.cell(r,11,af.get('tipo','AFASTAMENTO'))
        if af.get('obs'): obs_.append(af['obs'])
    rs=cfgm(CFG.get('rescisoes'),nm) if nm else None
    if rs: ws.cell(r,6,D(rs))
    ob=cfgm(CFG.get('obs'),nm) if nm else None
    if ob: obs_.append(ob)
    ws.cell(r,8,' '.join(obs_) if obs_ else None).font=f_norm
    ws.cell(r,8).alignment=Alignment(wrap_text=True,vertical='top')
ws.cell(3,4,'Férias, afastamento e rescisão informados aqui dispensam digitar ocorrência dia a dia: férias/afastamento não geram falta; dias após a rescisão não são apurados. Afastamento sem data de fim = em aberto.').font=f_note
for _n,_c in [('FUNC_NOME','A'),('FUNC_COD','B'),('FUNC_JORNADA','C'),('FUNC_FER_INI','D'),('FUNC_FER_FIM','E'),('FUNC_DESL','F'),('FUNC_CONTROLA','G'),('FUNC_AF_INI','I'),('FUNC_AF_FIM','J'),('FUNC_AF_TIPO','K')]:
    name(_n,f'FUNCIONARIOS!${_c}$5:${_c}${FL}')

# ================= PONTO_BRUTO =================
ws=ws_pb
title(ws,'PONTO_BRUTO - marcações de '+', '.join(os.path.basename(p) for p in CFG['espelhos'])+' (uma linha por funcionário/dia)','Colunas D a K = INPUT (edite aqui para tratar marcações; registre o motivo na col. K - Portaria 671 art.82/84). Colunas L a T (cinza) = valores do SOFTWARE do espelho, só para conferência.')
H=['Funcionário','Data','Dia','Ent.1','Saí.1','Ent.2','Saí.2','Ent.3','Saí.3','Observação / ocorrência','Tratamento manual (motivo)','CHPrev (sw)','Normais (sw)','Faltas (sw)','Atraso (sw)','Extras (sw)','ExSab (sw)','ExDom (sw)','Ad Not (sw)','Ex-Fa (sw, +/-)']
W=[30,11,6,8,8,8,8,8,8,22,30,9,9,9,9,9,9,9,9,10]
for i,h in enumerate(H,1): hdr(ws,4,i,h,sw=(i>=12),width=W[i-1])
ws.row_dimensions[4].height=32
def tm(m): return None if m is None else m/1440
r=FIRST
for x in recs:
    ws.cell(r,1,x['emp']); ws.cell(r,2,x['date']).number_format=T_DT; ws.cell(r,3,x['dow'])
    for k,m in enumerate(x['punches']): 
        c=ws.cell(r,4+k,tm(m)); c.number_format=T_CLK
    if x['note']: ws.cell(r,10,x['note'])
    if x.get('motivo'): ws.cell(r,11,x['motivo'])
    sw=[x['chprev'],x['normais'],x['faltas'],x['atraso'],x['extras'],x['exsab'],x['exdom'],x['adnot']]
    for k,m in enumerate(sw):
        if m: ws.cell(r,12+k,tm(m)).number_format=T_HM
    if x['exfa'] is not None and x['exfa']!=0: ws.cell(r,20,tm(x['exfa'])).number_format='+[h]:mm;-[h]:mm;'
    for c in range(1,21):
        cc=ws.cell(r,c); cc.font=f_in if 4<=c<=11 else f_norm; cc.border=box
        if 4<=c<=11: cc.fill=PatternFill('solid',fgColor='FFF2CC')
    r+=1
ws.freeze_panes='D5'; ws.auto_filter.ref=f'A4:T{LAST}'

# ================= CALCULO_DIARIO =================
ws=ws_ca
title(ws,'CALCULO_DIARIO - motor de cálculo (100% fórmulas; não digitar aqui)','★ laranja = fórmulas críticas de manutenção. Uma linha por linha do PONTO_BRUTO (mesma numeração de linhas). Horas em [h]:mm.')
cols=[  # (letra, título, largura, crit, formato, doc-regra)
 ('A','Funcionário',28,0,None,'Nome vindo do PONTO_BRUTO.'),
 ('B','Data',11,0,T_DT,'Data vinda do PONTO_BRUTO.'),
 ('C','Dia',6,0,None,'Dia da semana calculado da data.'),
 ('D','Tipo do dia',10,0,None,'Feriado (TABELAS) > Domingo > Sábado > Útil.'),
 ('E','Ocorrência',20,0,None,'Texto da col. J do PONTO_BRUTO (ATESTADO...) ou, se vazio, FÉRIAS automática quando a data cai no período informado em FUNCIONARIOS.'),
 ('F','Grupo',12,0,None,'Busca a ocorrência em TABELAS; "NÃO CADASTRADA" gera alerta.'),
 ('G','Abona?',8,0,None,'SIM = a ocorrência abona faltas (TABELAS col. C).'),
 ('H','Apurado (1/0)',8,0,'0','1 se a data está entre P_DT_INI e P_DT_FIM, não é posterior à rescisão (FUNCIONARIOS col. F) e o funcionário controla ponto (col. G).'),
 ('I','Jornada prevista',10,1,T_HM,'Dia útil = jornada do funcionário (ou P_JORNADA); sáb/dom/feriado = 0; fora do período = 0.'),
 ('J','Nº marcações',8,0,'0','Quantidade de horários digitados.'),
 ('K','Status',18,1,None,'FORA DO PERÍODO / OCORRÊNCIA / SEM MARCAÇÃO / FOLGA / MARCAÇÃO ÍMPAR / OK.'),
 ('L','Bruto (pares)',10,1,T_HM,'Soma (Saída-Entrada) dos pares; MOD(,1) trata virada de meia-noite. Zero se marcação ímpar.'),
 ('M','Intervalo deduzido',11,1,T_HM,'Só se P_INTERV_MODO=1: com 1 par e bruto > P_INTERV_LIM, desconta P_INTERV_PRE. No modo 2 (padrão) fica zero: o tempo é pago como extra.'),
 ('N','Trabalhado líquido',10,1,T_HM,'Bruto - intervalo pré-assinalado.'),
 ('O','Apurável (1/0)',8,0,'0','1 se o dia pode ser calculado (status OK, OCORRÊNCIA ou FOLGA). Marcação ímpar e dia útil sem marcação/observação = pendente (perguntar ao cliente).'),
 ('P','Saldo do dia',10,0,'+[h]:mm;-[h]:mm;','Trabalhado - Previsto (arredondado ao minuto).'),
 ('Q','Tolerância (1/0)',8,0,'0','1 se |saldo| <= P_TOL (CLT art.58 §1º): não gera falta nem extra.'),
 ('R','Faltas/atrasos (h)',10,1,T_HM,'Saldo negativo, fora da tolerância, sem ocorrência abonadora.'),
 ('S','Abonado (h)',10,0,T_HM,'Saldo negativo coberto por ocorrência que abona (informativo).'),
 ('T','Extras (h)',10,1,T_HM,'Saldo positivo fora da tolerância. Em sábado/domingo/feriado toda hora trabalhada é extra.'),
 ('U','Extras dia útil',10,0,T_HM,'Parte de T em dia útil.'),
 ('V','Extras sábado',10,0,T_HM,'Parte de T em sábado.'),
 ('W','Extras dom./feriado',10,0,T_HM,'Parte de T em domingo/feriado (100% - CCT cl.7 "c").'),
 ('X','Noturno real',10,1,T_HM,'Horas efetivamente trabalhadas entre P_NOT_INI e P_NOT_FIM (soma dos 3 pares).'),
 ('Y','Noturno c/ redução (52m30s)',11,1,T_HM,'X x (1h / P_HORA_NOT). Ex.: 48 min reais = 00:55.'),
 ('Z','Extras noturnas',10,0,T_HM,'MIN(Noturno real; Extras) - informativo (coluna "Ex Not" do software).'),
 ('AA','Falta integral (1/0)',8,1,'0','Dia com jornada prevista, sem horas trabalhadas e sem abono => 1 (vira Registro 11 no Domínio).'),
 ('AB','Semana (segunda)',11,0,T_DT,'Data da segunda-feira da semana (chave semanal).'),
 ('AC','1ª marcação',9,0,T_CLK,'Primeiro horário do dia.'),
 ('AD','Última marcação',9,0,T_CLK,'Último horário digitado do dia.'),
 ('AE','Interjornada',10,0,T_HM,'Da última marcação do dia anterior até a 1ª de hoje (só dias consecutivos completos).'),
 ('AF','Alertas',60,1,None,'Texto com todas as inconsistências do dia. Conferir ANTES de exportar.'),
 ('AG','Func. liberado?',9,0,None,'Lê status do funcionário no RESUMO_MENSAL (BLOQUEADO = fora da exportação).'),
 ('AH','Exporta falta (1/0)',8,0,'0','Falta integral de funcionário liberado.'),
 ('AI','Seq. falta',8,0,'0','Numeração corrida das faltas exportáveis (alimenta EXPORT_CALC).'),
 ('AJ','Pendência (1/0)',8,0,'0','1 se o dia precisa de confirmação com a empresa: marcação ímpar, dia útil sem marcação/ocorrência, ocorrência não cadastrada ou marcação em férias/afastamento.'),
 ('AK','Seq. pendência',8,0,'0','Numeração corrida das pendências (alimenta a aba PENDENCIAS).'),
 ('AL','Intervalo registrado',10,0,T_HM,'Só dias com 4 marcações: Ent.2 - Saí.1.'),
 ('AM','Jornada > 10h (1/0)',8,0,'0','CLT art.59: dia com trabalho líquido acima de P_LIM_DIA.'),
 ('AN','Interjornada < 11h (1/0)',8,0,'0','CLT art.66.'),
 ('AO','Almoço não registrado (1/0)',8,0,'0','Dia com 1 par de marcações e jornada acima de P_INTERV_LIM (CLT art.71); pela regra da empresa, vira hora extra.'),
 ('AP','Intervalo < 1h (1/0)',8,0,'0','CLT art.71 (verificar norma coletiva).'),
 ('AQ','Intervalo > 2h (1/0)',8,0,'0','CLT art.71.'),
 ('AR','Trabalho em férias/afast. (1/0)',8,0,'0','Marcação em dia de férias, afastamento ou licença.'),
 ('AS','Irregularidade legal (1/0)',8,0,'0','1 se qualquer teste legal acima disparou no dia.'),
 ('AT','Seq. sem marcação',8,0,'0','Numeração dos dias úteis sem marcação e sem observação (alimenta a lista "Perguntar ao cliente" do PAINEL).'),
]
CA={c[0]:c for c in cols}
for i,c in enumerate(cols,1): hdr(ws,4,i,c[1],crit=bool(c[3]),width=c[2])
ws.row_dimensions[4].height=44
PB="PONTO_BRUTO!"
def nightpair(e,s):
    S=f"IF({s}<{e},{s}+1,{s})"
    return (f"IF(AND({e}<>\"\",{s}<>\"\"),MAX(0,MIN({S},1+P_NOT_FIM)-MAX({e},P_NOT_INI))+MAX(0,MIN({S},P_NOT_FIM)-{e}),0)")
def pairdur(e,s): return f"IF(AND({e}<>\"\",{s}<>\"\"),MOD({s}-{e},1),0)"
def frm(r):
    p=lambda col: f"{PB}{col}{r}"
    d=lambda col: f"{col}{r}"
    pairs=[('D','E'),('F','G'),('H','I')]
    F={}
    F['A']=f"={p('A')}"; F['B']=f"={p('B')}"
    F['C']=f'=CHOOSE(WEEKDAY(B{r},2),"Seg","Ter","Qua","Qui","Sex","Sáb","Dom")'
    F['D']=f'=IF(COUNTIF(TB_FERIADOS,B{r})>0,"Feriado",IF(WEEKDAY(B{r},2)=7,"Domingo",IF(WEEKDAY(B{r},2)=6,"Sábado","Útil")))'
    _m=f'MATCH(A{r},FUNC_NOME,0)'
    F['E']=(f'=IF({p("J")}<>"",UPPER(TRIM({p("J")})),IFERROR(IF(AND(INDEX(FUNC_FER_INI,{_m})>0,B{r}>=INDEX(FUNC_FER_INI,{_m}),B{r}<=INDEX(FUNC_FER_FIM,{_m})),"FÉRIAS",'
            f'IF(AND(INDEX(FUNC_AF_INI,{_m})>0,B{r}>=INDEX(FUNC_AF_INI,{_m}),OR(INDEX(FUNC_AF_FIM,{_m})=0,B{r}<=INDEX(FUNC_AF_FIM,{_m}))),IF(INDEX(FUNC_AF_TIPO,{_m})="","AFASTAMENTO",UPPER(INDEX(FUNC_AF_TIPO,{_m}))),"")),""))')
    F['F']=f'=IF(E{r}="","",IFERROR(INDEX(TB_OC_GRUPO,MATCH(E{r},TB_OC_NOME,0)),"NÃO CADASTRADA"))'
    F['G']=f'=IF(E{r}="","NÃO",IFERROR(INDEX(TB_OC_ABONA,MATCH(E{r},TB_OC_NOME,0)),"NÃO"))'
    F['H']=f'=IF(AND(B{r}>=P_DT_INI,B{r}<=P_DT_FIM,IFERROR(OR(INDEX(FUNC_DESL,MATCH(A{r},FUNC_NOME,0))=0,B{r}<=INDEX(FUNC_DESL,MATCH(A{r},FUNC_NOME,0))),TRUE),IFERROR(INDEX(FUNC_CONTROLA,MATCH(A{r},FUNC_NOME,0))<>"NÃO",TRUE)),1,0)'
    F['I']=f'=IF(H{r}=0,0,IF(D{r}="Útil",IFERROR(INDEX(FUNC_JORNADA,MATCH(A{r},FUNC_NOME,0)),P_JORNADA),0))'
    F['J']=f'=COUNT({PB}D{r}:I{r})'
    F['K']=f'=IF(H{r}=0,"FORA DO PERÍODO",IF(J{r}=0,IF(E{r}<>"","OCORRÊNCIA",IF(I{r}>0,"SEM MARCAÇÃO","FOLGA")),IF(MOD(J{r},2)=1,"MARCAÇÃO ÍMPAR","OK")))'
    F['L']=f'=IF(MOD(J{r},2)=1,0,ROUND(('+'+'.join(pairdur(p(a+''),p(b)) for a,b in pairs)+')*1440,0)/1440)'
    F['M']=f'=IF(AND(P_INTERV_MODO=1,J{r}=2,L{r}>P_INTERV_LIM),P_INTERV_PRE,0)'
    F['N']=f'=MAX(0,L{r}-M{r})'
    F['O']=f'=IF(OR(K{r}="OK",K{r}="OCORRÊNCIA",K{r}="FOLGA"),1,0)'
    F['P']=f'=IF(O{r}=1,ROUND((N{r}-I{r})*1440,0)/1440,0)'
    F['Q']=f'=IF(AND(O{r}=1,P{r}<>0,ROUND(ABS(P{r})*1440,0)<=ROUND(P_TOL*1440,0)),1,0)'
    F['R']=f'=IF(AND(O{r}=1,P{r}<0,Q{r}=0,G{r}<>"SIM"),-P{r},0)'
    F['S']=f'=IF(AND(O{r}=1,P{r}<0,G{r}="SIM"),-P{r},0)'
    F['T']=f'=IF(AND(O{r}=1,P{r}>0,Q{r}=0),P{r},0)'
    F['U']=f'=IF(D{r}="Útil",T{r},0)'; F['V']=f'=IF(D{r}="Sábado",T{r},0)'; F['W']=f'=IF(OR(D{r}="Domingo",D{r}="Feriado"),T{r},0)'
    F['X']=f'=IF(AND(O{r}=1,J{r}>0,MOD(J{r},2)=0),ROUND(('+'+'.join(nightpair(p(a),p(b)) for a,b in pairs)+')*1440,0)/1440,0)'
    F['Y']=f'=ROUND(X{r}*(1/24)/P_HORA_NOT*1440,0)/1440'
    F['Z']=f'=MIN(X{r},T{r})'
    F['AA']=f'=IF(AND(O{r}=1,I{r}>0,N{r}=0,G{r}<>"SIM"),1,0)'
    F['AB']=f'=B{r}-WEEKDAY(B{r},2)+1'
    F['AC']=f'=IF(J{r}>0,{p("D")},"")'
    F['AD']=f'=IF(J{r}=0,"",IF({p("I")}<>"",{p("I")},IF({p("H")}<>"",{p("H")},IF({p("G")}<>"",{p("G")},IF({p("F")}<>"",{p("F")},IF({p("E")}<>"",{p("E")},{p("D")}))))))'
    F['AE']=f'=IFERROR(IF(AND(A{r}=A{r-1},B{r}-B{r-1}=1,AC{r}<>"",AD{r-1}<>"",K{r-1}="OK"),ROUND(((B{r}-B{r-1})+AC{r}-AD{r-1})*1440,0)/1440,""),"")'
    F['AF']=('=TRIM(IF(K{r}="MARCAÇÃO ÍMPAR","Marcação ímpar: corrigir antes de exportar. ","")&IF(K{r}="SEM MARCAÇÃO","Dia útil sem marcação e sem observação: perguntar ao cliente (falta, folga, férias ou atestado?). ","")'
             '&IF(AND(J{r}=2,L{r}>P_INTERV_LIM),IF(P_INTERV_MODO=1,"Só 1 par: intervalo pré-assinalado (deduzido). ","Só 1 par: almoço não registrado, pago como extra. "),"")&IF(N{r}>P_LIM_DIA,"Jornada > 10h (CLT 59). ","")'
             '&IF(AND(AE{r}<>"",N(AE{r})<P_INTERJ,AE{r}<>""),"Interjornada < 11h (CLT 66). ","")'
             '&IF(AND(J{r}>0,OR(F{r}="FÉRIAS",F{r}="AFASTAMENTO",F{r}="LICENÇA")),"Trabalho em férias/afastamento/licença. ","")'
             '&IF(F{r}="NÃO CADASTRADA","Ocorrência não cadastrada em TABELAS. ","")&IF(AP{r}=1,"Intervalo < 1h (CLT 71). ","")&IF(AQ{r}=1,"Intervalo > 2h (CLT 71). ","")&IF(X{r}>0,"Horas noturnas. ","")'
             '&IF(AND(D{r}<>"Útil",J{r}>0,O{r}=1),"Trabalho em "&LOWER(D{r})&". ",""))').format(r=r)
    F['AG']=f'=IFERROR(IF(OR(LEFT(INDEX(RESUMO_STATUS,MATCH(A{r},RESUMO_NOME,0)),9)="BLOQUEADO",LEFT(INDEX(RESUMO_STATUS,MATCH(A{r},RESUMO_NOME,0)),8)="IGNORADO"),"NÃO","SIM"),"NÃO")'
    F['AH']=f'=IF(AND(AA{r}=1,AG{r}="SIM"),1,0)'
    F['AI']=f'=IF(AH{r}=1,SUM(AH${FIRST}:AH{r}),"")'
    F['AJ']=f'=IF(AND(H{r}=1,OR(K{r}="MARCAÇÃO ÍMPAR",K{r}="SEM MARCAÇÃO",F{r}="NÃO CADASTRADA",AND(J{r}>0,OR(F{r}="FÉRIAS",F{r}="AFASTAMENTO",F{r}="LICENÇA")))),1,0)'
    F['AK']=f'=IF(AJ{r}=1,SUM(AJ${FIRST}:AJ{r}),"")'
    F['AL']=f'=IF(AND(O{r}=1,J{r}=4),ROUND(({p("F")}-{p("E")})*1440,0)/1440,"")'
    F['AM']=f'=IF(AND(O{r}=1,ROUND(N{r}*1440,0)>ROUND(P_LIM_DIA*1440,0)),1,0)'
    F['AN']=f'=IF(AND(AE{r}<>"",ROUND(N(AE{r})*1440,0)<ROUND(P_INTERJ*1440,0)),1,0)'
    F['AO']=f'=IF(AND(O{r}=1,J{r}=2,L{r}>P_INTERV_LIM),1,0)'
    F['AP']=f'=IF(AND(AL{r}<>"",ROUND(N(AL{r})*1440,0)<ROUND(P_INTERV_MIN*1440,0)),1,0)'
    F['AQ']=f'=IF(AND(AL{r}<>"",ROUND(N(AL{r})*1440,0)>ROUND(P_INTERV_MAX*1440,0)),1,0)'
    F['AR']=f'=IF(AND(H{r}=1,J{r}>0,OR(F{r}="FÉRIAS",F{r}="AFASTAMENTO",F{r}="LICENÇA")),1,0)'
    F['AS']=f'=IF(AM{r}+AN{r}+AO{r}+AP{r}+AQ{r}+AR{r}>0,1,0)'
    F['AT']=f'=IF(K{r}="SEM MARCAÇÃO",COUNTIF(K${FIRST}:K{r},"SEM MARCAÇÃO"),"")'
    return F
for r in range(FIRST,LAST+1):
    F=frm(r)
    for col,c in CA.items():
        cell=ws[f'{col}{r}']; cell.value=F[col]; cell.font=f_link if col in('A','B') else f_norm; cell.border=box
        if c[4]: cell.number_format=c[4]
ws.freeze_panes='C5'; ws.auto_filter.ref=f'A4:AT{LAST}'
rngK=f'K{FIRST}:K{LAST}'
ws.conditional_formatting.add(rngK,CellIsRule(operator='equal',formula=['"MARCAÇÃO ÍMPAR"'],fill=PatternFill('solid',bgColor='FF9999',fgColor='FF9999')))
ws.conditional_formatting.add(rngK,CellIsRule(operator='equal',formula=['"SEM MARCAÇÃO"'],fill=PatternFill('solid',bgColor='F8CBAD',fgColor='F8CBAD')))
ws.conditional_formatting.add(f'AF{FIRST}:AF{LAST}',FormulaRule(formula=[f'LEN(AF{FIRST})>0'],fill=PatternFill('solid',bgColor='FFF2CC',fgColor='FFF2CC')))
def cr(col): return f"CALCULO_DIARIO!${col}${FIRST}:${col}${LAST}"

# ================= SEMANAL =================
ws=ws_se
title(ws,'SEMANAL - verificação 44h/semana e perda de DSR (Lei 605/49 + CCT cl.27)','Semana = segunda a domingo. DSR perdido = falta integral injustificada na semana OU falta de pelo menos meio período (P_DSR_MEIO) em algum dia da semana.')
_w0=DT_INI-datetime.timedelta(days=DT_INI.weekday()); weeks=[_w0+datetime.timedelta(days=7*i) for i in range(6)]
SH=['Funcionário','Segunda','Domingo (DSR)','Trabalhado','Previsto','Extras apuradas','Faltas parciais','Faltas integrais (dias)','Semana > 44h?','Ajuste 44h (h)','DSR perdido (1/0)','Func. liberado?','Exporta DSR (1/0)','Seq. DSR','Aviso','Maior falta em 1 dia (h)']
SW=[28,11,12,11,10,10,10,10,9,10,9,9,9,7,48,12]
for i,h in enumerate(SH,1): hdr(ws,4,i,h,crit=(i in(9,10,11)),width=SW[i-1])
ws.row_dimensions[4].height=44
r=FIRST; SE_FIRST=r
for e in emps:
    for wk in weeks:
        ws.cell(r,1,e); ws.cell(r,2,wk)
        ws.cell(r,3,f'=B{r}+6')
        cr_=lambda col: cr(col)
        ws.cell(r,4,f'=SUMIFS({cr("N")},{cr("A")},$A{r},{cr("AB")},$B{r})')
        ws.cell(r,5,f'=SUMIFS({cr("I")},{cr("A")},$A{r},{cr("AB")},$B{r})')
        ws.cell(r,6,f'=SUMIFS({cr("T")},{cr("A")},$A{r},{cr("AB")},$B{r})')
        ws.cell(r,7,f'=SUMIFS({cr("R")},{cr("A")},$A{r},{cr("AB")},$B{r},{cr("AA")},0)')
        ws.cell(r,8,f'=SUMIFS({cr("AA")},{cr("A")},$A{r},{cr("AB")},$B{r})')
        ws.cell(r,9,f'=IF(ROUND(D{r}*1440,0)>ROUND(P_LIM_SEM*1440,0),"SIM","NÃO")')
        ws.cell(r,10,f'=MAX(0,ROUND((D{r}-P_LIM_SEM-F{r})*1440,0)/1440)')
        ws.cell(r,11,f'=IF(AND(C{r}>=P_DT_INI,C{r}<EDATE(P_DT_INI,1),OR(H{r}>0,ROUND(P{r}*1440,0)>=ROUND(P_DSR_MEIO*1440,0)-0.5)),1,0)')
        ws.cell(r,16,f'=SUMPRODUCT(MAX(({cr("A")}=$A{r})*({cr("AB")}=$B{r})*{cr("R")}))')
        ws.cell(r,12,f'=IFERROR(IF(OR(LEFT(INDEX(RESUMO_STATUS,MATCH(A{r},RESUMO_NOME,0)),9)="BLOQUEADO",LEFT(INDEX(RESUMO_STATUS,MATCH(A{r},RESUMO_NOME,0)),8)="IGNORADO"),"NÃO","SIM"),"NÃO")')
        ws.cell(r,13,f'=IF(AND(K{r}=1,L{r}="SIM"),1,0)')
        ws.cell(r,14,f'=IF(M{r}=1,SUM(M${FIRST}:M{r}),"")')
        ws.cell(r,15,f'=IF(B{r}<P_DT_INI,"Semana iniciada no período anterior: faltas antes de "&TEXT(P_DT_INI,"dd/mm")&" não estão neste espelho.",IF(I{r}="SIM","Semana acima de 44h.",""))')
        for c in range(1,17):
            cc=ws.cell(r,c); cc.font=f_norm; cc.border=box
        for c in (2,3): ws.cell(r,c).number_format=T_DT
        for c in (4,5,6,7,10,16): ws.cell(r,c).number_format=T_HM
        r+=1
SE_LAST=r-1
ws.freeze_panes='D5'; ws.auto_filter.ref=f'A4:P{SE_LAST}'
def se(col): return f"SEMANAL!${col}${SE_FIRST}:${col}${SE_LAST}"

# ================= RESUMO_MENSAL =================
ws=ws_re
title(ws,'RESUMO_MENSAL - apuração por funcionário (alimenta a exportação)','★ = colunas que viram rubricas no Domínio. Status BLOQUEADO = sem registro/ocorrência no mês (NÃO exportado até resolver). IGNORADO = não controla ponto (FUNCIONARIOS col. G).')
RH=[('Funcionário',28,0,None),('Cód. Domínio',10,0,'0'),('Dias c/ marcação',9,0,'0'),('Dias c/ marcação ímpar (pendentes)',11,0,'0'),('Faltas integrais (dias)',9,1,'0'),('Dias abonados (prev.>0)',9,0,'0'),
    ('Trabalhado',11,0,T_HM),('Previsto',11,0,T_HM),('Faltas parciais/atrasos (h)',11,0,T_HM),('Faltas integrais (h)',11,0,T_HM),('Abonos (h)',10,0,T_HM),
    ('Extras a 70% (dia útil; sáb. se P_SAB100=NÃO)',12,0,T_HM),('Extras a 100% direto (dom./fer.; sáb. se P_SAB100=SIM)',12,0,T_HM),('Noturno real',10,0,T_HM),('Noturno c/ redução',10,0,T_HM),('DSR perdidos (dias)',9,1,'0'),
    ('Faltas parciais a descontar (h)',11,1,T_HM),('Extras úteis/sáb líquidas',11,0,T_HM),('HE 70% (h)',10,1,T_HM),('HE 100% (h)',10,1,T_HM),('Adicional noturno (h)',10,1,T_HM),('Redução hora noturna (h)',10,1,T_HM),
    ('Saldo Ex-Fa (informativo)',11,0,'+[h]:mm;-[h]:mm;0:00'),('Dias com alerta',8,0,'0'),('Status',34,1,None),('Horas a descontar: falta (parcial + dias inteiros)',13,1,T_HM),('Horas de DSR a descontar',11,1,T_HM)]
for i,(h,w,cr_,fm) in enumerate(RH,1): hdr(ws,4,i,h,crit=bool(cr_),width=w)
ws.row_dimensions[4].height=58
RE_FIRST=5
for i,e in enumerate(emps):
    r=RE_FIRST+i
    ws.cell(r,1,e).font=f_link
    A=f'$A{r}'
    F={
     2:f'=IFERROR(INDEX(FUNC_COD,MATCH(A{r},FUNC_NOME,0)),"")',
     3:f'=COUNTIFS({cr("A")},{A},{cr("J")},">0")',
     4:f'=COUNTIFS({cr("A")},{A},{cr("K")},"MARCAÇÃO ÍMPAR")',
     5:f'=SUMIFS({cr("AA")},{cr("A")},{A})',
     6:f'=COUNTIFS({cr("A")},{A},{cr("G")},"SIM",{cr("I")},">0")',
     7:f'=SUMIFS({cr("N")},{cr("A")},{A})',
     8:f'=SUMIFS({cr("I")},{cr("A")},{A})',
     9:f'=SUMIFS({cr("R")},{cr("A")},{A},{cr("AA")},0)',
     10:f'=SUMIFS({cr("R")},{cr("A")},{A},{cr("AA")},1)',
     11:f'=SUMIFS({cr("S")},{cr("A")},{A})',
     12:f'=SUMIFS({cr("U")},{cr("A")},{A})+IF(P_SAB100="SIM",0,SUMIFS({cr("V")},{cr("A")},{A}))',
     13:f'=SUMIFS({cr("W")},{cr("A")},{A})+IF(P_SAB100="SIM",SUMIFS({cr("V")},{cr("A")},{A}),0)',
     14:f'=SUMIFS({cr("X")},{cr("A")},{A})',
     15:f'=SUMIFS({cr("Y")},{cr("A")},{A})',
     16:f'=SUMIFS({se("K")},{se("A")},{A})',
     17:f'=IF(P_MODO=2,MAX(0,I{r}-L{r}),I{r})',
     18:f'=IF(P_MODO=2,MAX(0,L{r}-I{r}),L{r})',
     19:f'=MIN(R{r},P_HE_LIM)',
     20:f'=MAX(0,R{r}-P_HE_LIM)+M{r}',
     21:f'=N{r}',
     22:f'=MAX(0,O{r}-N{r})',
     23:f'=(L{r}+M{r})-(I{r}+J{r})',
     24:f'=SUMPRODUCT(({cr("A")}={A})*(LEN({cr("AF")})>0))',
     26:f'=Q{r}+E{r}*P_HORAS_DIA_FALTA',
     27:f'=P{r}*P_HORAS_DIA_FALTA',
     25:f'=IF(IFERROR(INDEX(FUNC_CONTROLA,MATCH(A{r},FUNC_NOME,0)),"SIM")="NÃO","IGNORADO - não controla ponto",IF(AND(C{r}=0,SUMPRODUCT(({cr("A")}={A})*({cr("E")}<>""))=0),"BLOQUEADO - sem registro no mês",IF(SUMIFS({cr("AJ")},{cr("A")},$A{r})>0,"REVISAR - itens a confirmar","OK")))',
    }
    for c,fx in F.items():
        cell=ws.cell(r,c,fx); cell.font=f_norm; cell.border=box
        if RH[c-1][3]: cell.number_format=RH[c-1][3]
    ws.cell(r,1).border=box
RE_LAST=RE_FIRST+NE-1
name('RESUMO_NOME',f'RESUMO_MENSAL!$A${RE_FIRST}:$A${RE_LAST}'); name('RESUMO_STATUS',f'RESUMO_MENSAL!$Y${RE_FIRST}:$Y${RE_LAST}')
rt=RE_LAST+1
ws.cell(rt,1,'TOTAL').font=f_bold
for c in list(range(3,25))+[26,27]:
    if c in (2,): continue
    cell=ws.cell(rt,c,f'=SUM({L(c)}{RE_FIRST}:{L(c)}{RE_LAST})'); cell.font=f_bold; cell.border=box; cell.fill=fill_sub
    if RH[c-1][3]: cell.number_format=RH[c-1][3]
ws.cell(rt+2,1,'Pendências: cada dia com marcação ímpar (col. D) NÃO é calculado nem exportado até ser tratado no PONTO_BRUTO. Bloqueados: resolver a situação (férias? afastamento? admissão/demissão? falta de importação) e registrar a ocorrência.').font=f_note
ws.freeze_panes='B5'
ws.conditional_formatting.add(f'Y{RE_FIRST}:Y{RE_LAST}',FormulaRule(formula=[f'LEFT(Y{RE_FIRST},9)="BLOQUEADO"'],fill=PatternFill('solid',bgColor='FF9999',fgColor='FF9999')))
ws.conditional_formatting.add(f'Y{RE_FIRST}:Y{RE_LAST}',FormulaRule(formula=[f'LEFT(Y{RE_FIRST},7)="REVISAR"'],fill=PatternFill('solid',bgColor='FFE699',fgColor='FFE699')))
ws.conditional_formatting.add(f'Y{RE_FIRST}:Y{RE_LAST}',FormulaRule(formula=[f'Y{RE_FIRST}="OK"'],fill=PatternFill('solid',bgColor='C6E0B4',fgColor='C6E0B4')))

# ================= CONFERENCIA =================
ws=ws_cf
title(ws,'CONFERENCIA - software do espelho x recálculo desta planilha','Mostra onde o resultado do software difere do recálculo. Diferença positiva = software apura MAIS que esta planilha.')
CH=['Funcionário','Faltas software','Faltas recalculadas','Dif. faltas','Extras software','Extras recalculadas','Dif. extras','Dias com divergência (>tolerância)','Dias com marcação ímpar','Dias c/ 1 par só (sem almoço)','Status']
CW=[28,11,11,11,11,11,11,13,11,11,34]
for i,h in enumerate(CH,1): hdr(ws,4,i,h,width=CW[i-1])
ws.row_dimensions[4].height=44
pb=lambda col: f"PONTO_BRUTO!${col}${FIRST}:${col}${LAST}"
for i,e in enumerate(emps):
    r=5+i
    ws.cell(r,1,e).font=f_link
    ws.cell(r,2,f'=SUMIFS({pb("N")},{pb("A")},$A{r})')
    ws.cell(r,3,f'=INDEX(RESUMO_MENSAL!$I${RE_FIRST}:$I${RE_LAST},MATCH($A{r},RESUMO_NOME,0))+INDEX(RESUMO_MENSAL!$J${RE_FIRST}:$J${RE_LAST},MATCH($A{r},RESUMO_NOME,0))')
    ws.cell(r,4,f'=B{r}-C{r}')
    ws.cell(r,5,f'=SUMIFS({pb("P")},{pb("A")},$A{r})')
    ws.cell(r,6,f'=INDEX(RESUMO_MENSAL!$L${RE_FIRST}:$L${RE_LAST},MATCH($A{r},RESUMO_NOME,0))+INDEX(RESUMO_MENSAL!$M${RE_FIRST}:$M${RE_LAST},MATCH($A{r},RESUMO_NOME,0))')
    ws.cell(r,7,f'=E{r}-F{r}')
    ws.cell(r,8,f'=SUMPRODUCT(({cr("A")}=$A{r})*(({cr("O")}=1)*(ABS({cr("T")}-{pb("P")})+ABS({cr("R")}-{pb("N")})>P_TOL)))')
    ws.cell(r,9,f'=COUNTIFS({cr("A")},$A{r},{cr("K")},"MARCAÇÃO ÍMPAR")')
    ws.cell(r,10,f'=COUNTIFS({cr("A")},$A{r},{cr("M")},">0")')
    ws.cell(r,11,f'=INDEX(RESUMO_STATUS,MATCH($A{r},RESUMO_NOME,0))')
    for c in range(1,12):
        cc=ws.cell(r,c); cc.border=box
        if c>1: cc.font=f_norm
        if c in(2,3,4,5,6,7): cc.number_format='[h]:mm;-[h]:mm;0:00'
rt=5+NE
ws.cell(rt,1,'TOTAL').font=f_bold
for c in range(2,11):
    cell=ws.cell(rt,c,f'=SUM({L(c)}5:{L(c)}{rt-1})'); cell.font=f_bold; cell.fill=fill_sub; cell.border=box
    if c<=7: cell.number_format='[h]:mm;-[h]:mm;0:00'
ws.freeze_panes='B5'


# ================= parte 2: exportação, demonstração, dicionário, leia-me =================
TESTE=False   # abas de teste de agosto removidas (continham dados reais de funcionários)
# --- patch: código Domínio vazio deve ficar "" (não 0)
for i,e in enumerate(emps):
    r=RE_FIRST+i
    ws_re.cell(r,2).value=f'=IFERROR(IF(INDEX(FUNC_COD,MATCH(A{r},FUNC_NOME,0))="","",INDEX(FUNC_COD,MATCH(A{r},FUNC_NOME,0))),"")'

# ---------- EXPORT_CALC ----------
ws=ws_ec
title(ws,'EXPORT_CALC - montagem dos lançamentos para o Domínio (fórmulas; não digitar)','Bloco 1 = rubricas em horas (Registro 10). Bloco 2 = faltas integrais por data. Bloco 3 = DSR perdido por data. Bloco 4 = fila de saída que alimenta EXPORT_CSV / EXPORT_TXT.')
E1=5; NR=6; E1L=E1+NE*NR-1
rub=[('HE 70%','R_HE70',19),('HE 100% (>30h e dom./fer.)','R_HE100',20),('Adicional noturno (horas)','R_NOT',21),('Redução da hora noturna','R_NOTRED',22),('Faltas em horas (parciais + dias inteiros)','R_FALTA',26),('DSR perdido em horas','R_DSR',27)]
for i,h in enumerate(['Funcionário','Cód. empregado','Rubrica','Cód. rubrica','Horas (hh:mm)','Horas decimais','Status do funcionário','Incluir (1/0)','Seq.'],1):
    hdr(ws,4,i,h,crit=(i in(8,)),width=[28,11,30,10,11,10,32,8,6][i-1])
r=E1
for i,e in enumerate(emps):
    rr=RE_FIRST+i
    for (desc,rc,col) in rub:
        ws.cell(r,1,f'=RESUMO_MENSAL!$A${rr}'); ws.cell(r,2,f'=RESUMO_MENSAL!$B${rr}'); ws.cell(r,3,desc)
        ws.cell(r,4,f'=IF({rc}="","",{rc})')
        ws.cell(r,5,f'=RESUMO_MENSAL!${L(col)}${rr}').number_format=T_HM
        ws.cell(r,6,f'=ROUND(E{r}*24,2)').number_format='0.00'
        ws.cell(r,7,f'=RESUMO_MENSAL!$Y${rr}')
        extra=f',D{r}<>""' if rc=='R_NOTRED' else ''
        ws.cell(r,8,f'=IF(AND(ROUND(E{r}*1440,0)>0,LEFT(G{r},9)<>"BLOQUEADO",LEFT(G{r},8)<>"IGNORADO"{extra}),1,0)')
        ws.cell(r,9,f'=IF(H{r}=1,SUM(H${E1}:H{r}),"")')
        for c in range(1,10): ws.cell(r,c).font=f_norm; ws.cell(r,c).border=box
        r+=1
# bloco 2 faltas integrais
NF=80; ND=30
hdr(ws,4,11,'Seq.',width=6); hdr(ws,4,12,'Data da falta integral',width=13); hdr(ws,4,13,'Funcionário',width=28); hdr(ws,4,14,'Cód. empregado',width=11)
for k in range(1,NF+1):
    r=E1+k-1
    ws.cell(r,11,k)
    ws.cell(r,12,f'=IFERROR(INDEX({cr("B")},MATCH(K{r},{cr("AI")},0)),"")').number_format=T_DT
    ws.cell(r,13,f'=IFERROR(INDEX({cr("A")},MATCH(K{r},{cr("AI")},0)),"")')
    ws.cell(r,14,f'=IF(M{r}="","",IFERROR(IF(INDEX(FUNC_COD,MATCH(M{r},FUNC_NOME,0))="","",INDEX(FUNC_COD,MATCH(M{r},FUNC_NOME,0))),""))')
    for c in range(11,15): ws.cell(r,c).font=f_norm; ws.cell(r,c).border=box
# bloco 3 DSR
hdr(ws,4,16,'Seq.',width=6); hdr(ws,4,17,'Domingo (DSR perdido)',width=13); hdr(ws,4,18,'Funcionário',width=28); hdr(ws,4,19,'Cód. empregado',width=11)
for k in range(1,ND+1):
    r=E1+k-1
    ws.cell(r,16,k)
    ws.cell(r,17,f'=IFERROR(INDEX({se("C")},MATCH(P{r},{se("N")},0)),"")').number_format=T_DT
    ws.cell(r,18,f'=IFERROR(INDEX({se("A")},MATCH(P{r},{se("N")},0)),"")')
    ws.cell(r,19,f'=IF(R{r}="","",IFERROR(IF(INDEX(FUNC_COD,MATCH(R{r},FUNC_NOME,0))="","",INDEX(FUNC_COD,MATCH(R{r},FUNC_NOME,0))),""))')
    for c in range(16,20): ws.cell(r,c).font=f_norm; ws.cell(r,c).border=box
# contadores e painel de pendências (colunas U:V)
hdr(ws,4,21,'Painel de pré-exportação',width=46); hdr(ws,4,22,'Valor',width=14)
b1=f'$H${E1}:$H${E1L}'
panel=[
 ('Linhas Registro 10 (rubricas em horas)',f'=COUNTIF({b1},1)'),
 ('(opcional) Faltas integrais por data - Registro 11',f'=COUNT($L${E1}:$L${E1+NF-1})'),
 ('(opcional) DSR perdidos por data - Registro 11',f'=COUNT($Q${E1}:$Q${E1+ND-1})'),
 ('Pendência: itens a confirmar com a empresa (lista no PAINEL)',f'=COUNT({cr("AK")})'),
 ('Pendência: funcionários BLOQUEADOS (sem registro)',f'=SUMPRODUCT(--(LEFT(RESUMO_STATUS,9)="BLOQUEADO"))'),
 ('Pendência: códigos não preenchidos (empresa/rubrica/empregado)',f'=IF(P_EMPRESA="",1,0)+SUMPRODUCT(({b1}=1)*((($B${E1}:$B${E1L}="")+($D${E1}:$D${E1L}=""))>0))+SUMPRODUCT(($L${E1}:$L${E1+NF-1}<>"")*($N${E1}:$N${E1+NF-1}=""))+SUMPRODUCT(($Q${E1}:$Q${E1+ND-1}<>"")*($S${E1}:$S${E1+ND-1}=""))'),
 ('SITUAÇÃO DA EXPORTAÇÃO','=IF(SUM(V8:V10)=0,"PRONTO PARA IMPORTAR","NÃO IMPORTAR: resolver pendências")'),
]
for i,(a,b) in enumerate(panel):
    r=5+i
    ws.cell(r,21,a).font=f_bold if i==6 else f_norm; ws.cell(r,22,b).font=f_bold if i==6 else f_norm
    ws.cell(r,21).border=box; ws.cell(r,22).border=box
    ws.cell(r,21).alignment=Alignment(wrap_text=True)
ws.conditional_formatting.add('V11',FormulaRule(formula=['LEFT(V11,3)="NÃO"'],fill=PatternFill('solid',bgColor='FF9999',fgColor='FF9999')))
ws.conditional_formatting.add('V11',FormulaRule(formula=['LEFT(V11,3)="PRO"'],fill=PatternFill('solid',bgColor='C6E0B4',fgColor='C6E0B4')))
ws.cell(13,21,'Faltas de dia inteiro e DSR perdido entram no TXT como HORAS (P_HORAS_DIA_FALTA por dia). Os blocos 2 e 3 (por data) alimentam apenas as linhas opcionais Registro 11 do EXPORT_CSV.').font=f_note
ws.cell(13,21).alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[13].height=54
# bloco 4 fila de saída
QN=max(160,E1L-E1+1+NF+ND)
hdr(ws,4,24,'Fila k',width=6); hdr(ws,4,25,'Tipo (10/11/12)',width=8); hdr(ws,4,26,'Posição no bloco',width=9)
for k in range(1,QN+1):
    r=E1+k-1
    ws.cell(r,24,k)
    ws.cell(r,25,f'=IF(X{r}<=$V$5,10,IF(X{r}<=$V$5+$V$6,11,IF(X{r}<=$V$5+$V$6+$V$7,12,"")))')
    ws.cell(r,26,f'=IF(Y{r}=10,MATCH(X{r},$I${E1}:$I${E1L},0),IF(Y{r}=11,X{r}-$V$5,IF(Y{r}=12,X{r}-$V$5-$V$6,"")))')
    for c in range(24,27): ws.cell(r,c).font=f_norm; ws.cell(r,c).border=box
ws.freeze_panes='A5'
ECQ=lambda r: (f'EXPORT_CALC!$Y${E1+r-2}',f'EXPORT_CALC!$Z${E1+r-2}')
def blk(col,rng_first,rng_last): return f'EXPORT_CALC!${col}${rng_first}:${col}${rng_last}'

# ---------- EXPORT_CSV ----------
ws=ws_csv
CH=['Registro','Empresa','Cod_empregado','Nome','Competencia','Cod_rubrica','Descricao','Tipo_folha','Horas_decimal','Horas_hhmm','Data_falta','Tipo_falta']
for i,h in enumerate(CH,1):
    c=ws.cell(1,i,h); c.font=f_hdr; c.fill=fill_hdr
    ws.column_dimensions[L(i)].width=[9,9,14,30,12,12,34,10,13,11,12,9][i-1]
for k in range(1,QN+1):
    r=k+1; Y,Z=ECQ(r)
    ws.cell(r,1,f'=IF({Y}="","",IF({Y}=10,"10","11"))')
    ws.cell(r,2,f'=IF(OR({Y}="",P_EMPRESA=""),"",P_EMPRESA)')
    ws.cell(r,3,f'=IF({Y}="","",IF({Y}=10,INDEX({blk("B",E1,E1L)},{Z}),IF({Y}=11,INDEX({blk("N",E1,E1+NF-1)},{Z}),INDEX({blk("S",E1,E1+ND-1)},{Z}))))')
    ws.cell(r,4,f'=IF({Y}="","",IF({Y}=10,INDEX({blk("A",E1,E1L)},{Z}),IF({Y}=11,INDEX({blk("M",E1,E1+NF-1)},{Z}),INDEX({blk("R",E1,E1+ND-1)},{Z}))))')
    ws.cell(r,5,f'=IF({Y}="","",P_COMP)')
    ws.cell(r,6,f'=IF({Y}=10,INDEX({blk("D",E1,E1L)},{Z}),"")')
    ws.cell(r,7,f'=IF({Y}="","",IF({Y}=10,INDEX({blk("C",E1,E1L)},{Z}),IF({Y}=11,"Falta integral (dia)","DSR perdido (falta na semana)")))')
    ws.cell(r,8,f'=IF({Y}="","",P_TIPOFOLHA)')
    ws.cell(r,9,f'=IF({Y}=10,INDEX({blk("F",E1,E1L)},{Z}),"")').number_format='0.00'
    ws.cell(r,10,f'=IF({Y}=10,INDEX({blk("E",E1,E1L)},{Z}),"")').number_format='[h]:mm'
    ws.cell(r,11,f'=IF({Y}="","",IF({Y}=10,"",IF({Y}=11,INDEX({blk("L",E1,E1+NF-1)},{Z}),INDEX({blk("Q",E1,E1+ND-1)},{Z}))))').number_format='dd/mm/yyyy'
    ws.cell(r,12,f'=IF({Y}="","",IF({Y}=10,"",IF({Y}=11,1,2)))')
    for c in range(1,13): ws.cell(r,c).font=f_norm
ws.freeze_panes='A2'

# ---------- EXPORT_TXT ----------
ws=ws_txt
ws.column_dimensions['A'].width=60
for k in range(1,QN+1):
    r=k; Y,Z=ECQ(r+1)
    cod=f'INDEX({blk("B",E1,E1L)},{Z})'; rb=f'INDEX({blk("D",E1,E1L)},{Z})'; hh=f'INDEX({blk("F",E1,E1L)},{Z})'
    ws.cell(r,1,f'=IF({Y}<>10,"",IF(OR({cod}="",{rb}="",P_EMPRESA=""),"ERRO-PREENCHER-CODIGOS","10"&TEXT({cod},"0000000000")&P_COMP&TEXT({rb},"0000")&TEXT(P_TIPOFOLHA,"00")&TEXT(ROUND({hh}*P_FATOR,0),"000000000")&TEXT(P_EMPRESA,"0000000000")))')
    ws.cell(r,1).font=Font(name='Courier New',size=10)

# ---------- DICIONARIO ----------
ws=ws_dic
title(ws,'DICIONARIO - o que cada coluna calcula e qual fórmula usa','★ = fórmula crítica (mexer só com cuidado). Fórmula mostrada = exemplo da linha 6 da aba. Parâmetros (P_...) ficam em PARAMETROS.')
for i,h in enumerate(['Aba','Coluna','Nome','Fórmula (exemplo)','Regra / o que faz','Crítica?'],1): hdr(ws,4,i,h,width=[16,8,28,90,70,9][i-1])
r=5
def put(aba,col,nome,fx,regra,crit):
    global r
    ws.cell(r,1,aba); ws.cell(r,2,col); ws.cell(r,3,nome)
    c=ws.cell(r,4,fx); c.data_type='s'; c.font=Font(name='Courier New',size=9)
    ws.cell(r,5,regra); ws.cell(r,6,'★' if crit else '')
    for k in (1,2,3,5,6): ws.cell(r,k).font=f_norm
    for k in (3,4,5): ws.cell(r,k).alignment=Alignment(wrap_text=True,vertical='top')
    for k in range(1,7): ws.cell(r,k).border=box
    r+=1
F6=frm(FIRST+1)
for col,c in CA.items(): put('CALCULO_DIARIO',col,c[1],F6[col],c[5],c[3])
fs=lambda r_: None
for i,(h,w,cr_,fm) in enumerate(RH,1):
    cell=ws_re.cell(RE_FIRST,i).value
    put('RESUMO_MENSAL',L(i),h,str(cell) if cell else '(nome do funcionário - vem do PONTO_BRUTO)','Soma/consulta por funcionário no CALCULO_DIARIO ou SEMANAL.',cr_)
for i,h in enumerate(SH,1):
    if i in (1,2): continue
    put('SEMANAL',L(i),h,str(ws_se.cell(SE_FIRST,i).value),'Chave = funcionário + segunda-feira da semana.',i in(9,10,11))
ws.freeze_panes='A5'


# ---------- PENDENCIAS ----------
ws_pd=wb.create_sheet('PENDENCIAS'); ws_pd.sheet_properties.tabColor='C00000'
ws=ws_pd
title(ws,'PENDENCIAS - itens para confirmar com a empresa antes de exportar','Lista automática: só aparece o que ainda está em aberto. Depois da resposta da empresa, corrija no PONTO_BRUTO (marcações D-I ou ocorrência J) e escreva o motivo na col. K: o item some sozinho e o dia é recalculado.')
ws['A3']=f'="Itens em aberto: "&COUNT({cr("AK")})'; ws['A3'].font=Font(name=FONT,size=11,bold=True,color='C00000')
PH=['Nº','Funcionário','Data','Dia','Marcações registradas','Tipo','O que confirmar com a empresa','Mensagem pronta para enviar','Linha PONTO_BRUTO']
PW=[5,30,11,6,26,26,60,95,9]
for i,h_ in enumerate(PH,1): hdr(ws,5,i,h_,width=PW[i-1])
ws.row_dimensions[5].height=32
NP=120
pbk=lambda col: f"PONTO_BRUTO!${col}${FIRST}:${col}${LAST}"
for k in range(1,NP+1):
    r=5+k
    ws.cell(r,1,k)
    ws.cell(r,9,f'=IFERROR(MATCH(A{r},{cr("AK")},0),"")')
    ws.cell(r,2,f'=IF(I{r}="","",INDEX({cr("A")},I{r}))')
    ws.cell(r,3,f'=IF(I{r}="","",INDEX({cr("B")},I{r}))').number_format=T_DT
    ws.cell(r,4,f'=IF(I{r}="","",INDEX({cr("C")},I{r}))')
    ws.cell(r,5,f'=IF(I{r}="","",TRIM('+'&'.join(f'IF(ISNUMBER(INDEX({pbk(c)},I{r})),TEXT(INDEX({pbk(c)},I{r}),"hh:mm")&" ","")' for c in 'DEFGHI')+'))')
    ws.cell(r,6,f'=IF(I{r}="","",IF(INDEX({cr("K")},I{r})="MARCAÇÃO ÍMPAR","Marcação ímpar",IF(INDEX({cr("K")},I{r})="SEM MARCAÇÃO","Dia útil sem marcação",IF(INDEX({cr("F")},I{r})="NÃO CADASTRADA","Ocorrência não cadastrada","Marcação em férias/afastamento"))))')
    ws.cell(r,7,f'=IF(I{r}="","",IF(F{r}="Marcação ímpar","Informar o horário que falta (ou confirmar esquecimento).",IF(F{r}="Dia útil sem marcação","Foi falta, folga, férias, atestado ou outra ocorrência?",IF(F{r}="Ocorrência não cadastrada","Cadastrar a ocorrência em TABELAS (abona ou não).","Férias/afastamento interrompido ou marcação indevida?"))))')
    ws.cell(r,8,f'=IF(I{r}="","",B{r}&" - "&TEXT(C{r},"dd/mm/yyyy")&" ("&D{r}&"): "&IF(E{r}="","","marcações "&E{r}&"- ")&G{r})')
    for c in range(1,10): ws.cell(r,c).font=f_norm; ws.cell(r,c).border=box
    for c in (7,8): ws.cell(r,c).alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A6'
ws.conditional_formatting.add(f'B6:H{5+NP}',FormulaRule(formula=['$B6<>""'],fill=PatternFill('solid',bgColor='FFF2CC',fgColor='FFF2CC')))
_t=wb.sheetnames.index('RESUMO_MENSAL')+1; wb.move_sheet('PENDENCIAS',offset=_t-wb.sheetnames.index('PENDENCIAS'))


# ---------- LEIA_ME ----------
ws=ws_lm
ws.column_dimensions['A'].width=5; ws.column_dimensions['B'].width=34; ws.column_dimensions['C'].width=120
ws['A1']=f'SISTEMA DE APURAÇÃO DE PONTO -> FOLHA DA DOMÍNIO  |  {EMPNOME.upper()}  |  competência {COMP[4:]}/{COMP[:4]}'; ws['A1'].font=f_title
ws['A2']='Automação sem macros: tudo é fórmula. Se um horário do PONTO_BRUTO ou um parâmetro mudar, o resultado, o resumo e a exportação recalculam sozinhos.'; ws['A2'].font=f_note
rr=4; _nsec=0
def sec(t):
    global rr,_nsec
    _nsec+=1; t=f'{_nsec}. '+re.sub(r'^\d+\.\s*','',t)   # numeração sempre sequencial
    c=ws.cell(rr,1,t); c.font=f_hdr; c.fill=fill_hdr
    for k in (2,3): ws.cell(rr,k).fill=fill_hdr
    rr+=1
def line(a,b,bold=False,fill=None):
    global rr
    ws.cell(rr,2,a).font=f_bold; c=ws.cell(rr,3,b); c.font=f_bold if bold else f_norm
    c.alignment=Alignment(wrap_text=True,vertical='top'); ws.cell(rr,2).alignment=Alignment(wrap_text=True,vertical='top')
    if fill:
        ws.cell(rr,2).fill=fill; ws.cell(rr,3).fill=fill
    n=max(1,len(b)//115+b.count('\n')+1); ws.row_dimensions[rr].height=max(15,13.5*n)
    rr+=1
h=lambda m: f"{m//60}:{m%60:02d}"
sec('1. COMO LER ESTE ARQUIVO')
line('PAINEL (1ª aba)','Visão geral do mês: dias sem marcação e sem observação (perguntar ao cliente), itens a confirmar, irregularidades legais, ocorrências e resultado por funcionário. O semáforo no topo diz se a exportação para o Domínio está liberada.')
line('Lista de pendências (PAINEL)','Itens que precisam de resposta do cliente ficam no topo do PAINEL. Clique no nome para ir à célula do PONTO_BRUTO a corrigir (marcações D-I ou ocorrência J + motivo na K): o item some e o dia é recalculado.')
line('Regras do cliente','Ficam em PARAMETROS, FUNCIONARIOS e TABELAS (jornada, almoço, tolerância, faixa de 30h, sábado 100%, DSR, rubricas e códigos do Domínio).')
sec('3. RECOMENDAÇÃO DE PLATAFORMA')
line('Escolha','EXCEL (.xlsx) com fórmulas puras, sem VBA - e o mesmo arquivo abre no Google Sheets se você quiser compartilhar.')
line('Por quê','(1) O Domínio Folha é desktop e importa arquivo TEXTO por Utilitários > Importação > De Arquivo Texto > De Lançamentos: gerar/salvar o TXT no Excel é direto. (2) Sem macro = sem alerta de segurança, sem manutenção de código, funciona no Excel 2016+, Google Sheets e LibreOffice. (3) Fórmulas (SUMIFS/INDEX/MATCH) auditáveis célula a célula. (4) Apps Script/VBA só valeriam para automatizar a carga mensal das abas do relógio; como o layout do espelho é irregular (uma aba por pessoa, colunas auxiliares no meio), a carga é mais segura feita por rotina de conversão validada - posso rodá-la aqui todo mês.')
line('CSV x TXT','O Domínio NÃO importa CSV puro: o leiaute é TXT posicional (manual "Importar Lançamentos": Registro 10 = lançamento de rubrica; 11 = data de falta; 12 = faltas parciais). Por isso entrego os dois: EXPORT_TXT (o que o Domínio lê) e EXPORT_CSV (conferência, arquivo de apoio ou conversão de planilhas do Domínio).')
sec('4. FLUXO MENSAL (passo a passo)')
line('Passo 1','Receber do cliente o espelho do relógio (Excel), a planilha-modelo preenchida, fotos ou PDFs do ponto. Dois cliques em "PONTO - 1 Gerar planilha do mes.bat" e escolha os arquivos: a planilha do mês é criada e aberta.')
line('Informações do RH junto com o espelho','A cada envio, informar: (a) funcionários em férias, com início e fim; (b) rescisões, com o último dia trabalhado; (c) quem não controla ponto (sócios/proprietários). São registrados em FUNCIONARIOS (colunas D a G): férias não geram falta, dias após a rescisão não são apurados e quem não controla ponto é ignorado. A jornada de 08:45 prevalece sobre o CHPrev (08:30) do espelho, que é ignorado.')
line('Passo 2','PONTO_BRUTO já vem carregado (D a K = marcações/ocorrência; L a T = valores do software do relógio, só conferência). Se precisar regerar a planilha (ex.: mudou o config), escolha a própria planilha corrigida como arquivo de entrada: as correções e motivos do PONTO_BRUTO são mantidos.')
line('Passo 3','Conferir PARAMETROS (datas, jornada, intervalo, tolerância, modo faltas x extras, códigos) e FUNCIONARIOS (código Domínio, períodos de férias e data de rescisão de quem tiver).')
line('Passo 4','Abrir o PAINEL (lista de pendências no topo: clique no nome para ir à célula a corrigir; seção A = fora da lei) e RESUMO_MENSAL (col. Status). Envie as mensagens à empresa; com as respostas, corrija o PONTO_BRUTO (marcações D-I ou ocorrência J + motivo na K). Trate também os alertas legais em CALCULO_DIARIO (col. Alertas).')
line('Passo 5','Abrir CONFERENCIA (software x recálculo) e o semáforo de exportação no topo do PAINEL. Só prossiga com "PRONTO PARA IMPORTAR".')
line('Passo 6 - TXT (recomendado)','Salve e feche a planilha. Dois cliques em "PONTO - 2 Gerar TXT para o Dominio.bat" e escolha a planilha: o programa recalcula, confere o semáforo e grava lancamentos_dominio_AAAAMM.txt (ANSI) na mesma pasta. No Domínio Folha: Utilitários > Importação > De Arquivo Texto > De Lançamentos > selecione o arquivo, competência, "Não importar" lançamentos existentes > Importar > Gravar. Confira em Processos > Lançamentos > Por empregado.')
line('Passo 6 - CSV','Aba EXPORT_CSV: Excel: Salvar como > CSV (separado por ponto e vírgula). Google Sheets: Arquivo > Fazer download > CSV (baixa só a aba ativa). Colunas na ordem: Registro; Empresa; Cod_empregado; Nome; Competencia; Cod_rubrica; Descricao; Tipo_folha; Horas_decimal; Horas_hhmm; Data_falta; Tipo_falta. Delete as linhas vazias ao final antes de salvar.')
line('Passo 7 - Faltas de dia inteiro e DSR','Entram no TXT como horas: falta de dia inteiro = P_HORAS_DIA_FALTA (07:20) por dia na rubrica 8069, e DSR perdido = 07:20 por semana na rubrica 42 (mesmo padrão dos meses anteriores no Domínio). As linhas Registro 11 (data + tipo 1/2) ficam no EXPORT_CSV como alternativa opcional, caso prefira lançar por data.')
sec('5. REGRAS APLICADAS (base) - onde customizar: PARAMETROS')
rules=[('Jornada','08:45 seg-sex (informado pelo RH; o CHPrev do espelho está em 08:30 e precisa ser corrigido no relógio/software). CCT cl.24: a prorrogação seg-sex compensa o sábado; esses minutos não são extras. Extra = horas acima de 08:45, não acima de 8:00. Se não houver compensação de sábado, mude P_JORNADA para 08:00.'),
 ('Limite 44h','Verificação semanal em SEMANAL (col. Semana > 44h / Ajuste 44h). Com 08:45 x 5 = 43:45 só estoura se houver sábado trabalhado.'),
 ('Horas extras','CCT cl.7: até 30h por mês a 70% e acima disso 100% (confirmado pelo RH); domingos e feriados = 100%. Sábado trabalhado = 100% direto: é dia de descanso, já compensado durante a semana (confirmado pelo RH; mais favorável que a CCT), e não conta na faixa de 30h. Extras de dia útil = tudo acima de 08:45.'),
 ('Adicional noturno','CCT cl.8: 22:00-05:00, 35%. Hora noturna reduzida 52m30s (CLT art.73 §1º) aplicada em "Noturno c/ redução".'),
 ('Tolerância','CLT art.58 §1º: até 10 min/dia não descontam nem geram extra (o software do espelho não aplica). P_TOL = 00:00 desliga.'),
 ('Faltas e DSR','Falta injustificada desconta as horas/dia. DSR: regra da empresa (informada pelo RH) = a partir de MEIO PERÍODO de falta no dia (metade da jornada, 04:22 - parâmetro P_DSR_MEIO) desconta o DSR da semana; falta integral sempre perde. Só conta falta injustificada: atestado (CCT cl.35), férias, licença e faltas justificadas (CCT cl.28) não descontam. Obs.: a Lei 605/49 art.6º e a CCT cl.27 são mais rigorosas com atrasos (só toleram até 30 min na semana); a regra da empresa é mais favorável ao empregado - é decisão da empresa.'),
 ('Faltas x extras','P_MODO=1 (padrão): paga extras e desconta faltas separadamente. P_MODO=2: compensa no mês (só saldo). O espelho mostra "Ex-Fa" (saldo), mas banco de horas exige acordo/consulta (CCT cl.29); não presumi.'),
 ('Fechamento','CCT cl.31 permite apurar de 21 a 20 ou 16 a 15; o espelho é do mês fechado 01-31/08. Ajuste P_DT_INI/P_DT_FIM se mudar a regra.'),
 ('Ainda a validar por profissional','Vínculo Registro 10/11 do Domínio. Esta planilha não substitui parecer do contador/jurídico da empresa.')]
rules.append(('Dia útil sem marcação e sem observação','Não é falta automática: aparece na lista de pendências do PAINEL. Se o cliente confirmar falta, registre FALTA INJUSTIFICADA na coluna J do PONTO_BRUTO (desconta o dia e o DSR); se for folga, férias ou atestado, registre a ocorrência correspondente.'))
for a,b in rules: line(a,b)
sec('6. MAPA DAS ABAS E FÓRMULAS CRÍTICAS (★ laranja no cabeçalho)')
for a,b in [x for x in [('PARAMETROS / FUNCIONARIOS / TABELAS','Entradas de regra (amarelo, fonte azul). Ninguém precisa mexer em fórmulas para mudar uma regra.'),
 ('PONTO_BRUTO','Dados do relógio: uma linha por funcionário/dia. Marcações (D-I), ocorrência (J), motivo do tratamento (K); cinza = valores do software só para conferência.'),
 ('CALCULO_DIARIO (oculta)','Motor do cálculo dia a dia - não excluir: RESUMO_MENSAL, PAINEL e exportação dependem dela. Para ver: clique direito em uma aba > Reexibir. Críticas: I (previsto), K (status), L (bruto), M (intervalo deduzido - só modo 1), N (líquido), R (faltas), T (extras), X/Y (noturno), AA (falta integral), AF (alertas).'),
 ('SEMANAL (oculta)','Verifica 44h/semana e decide a perda de DSR (meio período de falta). Críticas: I, J, K, P.'),
 ('RESUMO_MENSAL','Apuração por funcionário; aplica faixa 70%/100% e modo faltas x extras. Críticas: E, P, Q, S, T, U, V, Y (status).'),
 ('PAINEL','Dashboard mensal: KPIs, itens a confirmar com a empresa, irregularidades legais (jornada >10h, interjornada <11h, almoço, intervalo, trabalho em férias), ocorrências, qualidade do software do relógio, resultado por funcionário e gráficos. Abre por padrão.'),
 ('PENDENCIAS (oculta)','Base da lista de pendências do PAINEL (com mensagem pronta para copiar). Some quando o PONTO_BRUTO é corrigido.'),
 ('CONFERENCIA','Software x recálculo, por funcionário.'),
 ('EXPORT_CALC (oculta) / EXPORT_CSV / EXPORT_TXT','Montagem e saída. EXPORT_CALC!V11 = semáforo de exportação. TXT posicional Registro 10: "10" + cód. empregado(10) + AAAAMM + cód. rubrica(4) + tipo folha(2) + valor(9) + empresa(10) = 43 posições. Valor com 2 casas decimais implícitas (confirmado no manual); horas = horas decimais x 100. Teste com 1 empregado antes de importar tudo.'),
 ('DICIONARIO','Cada coluna, com a fórmula e a regra.'),
 ('Fontes','Convenção Coletiva 2026/2027 (SC000921/2026), Portaria MTP 671/2021 (arquivos do projeto); manual "Importar Lançamentos" do Domínio e relatório Movimentos (21/09/2026) enviados pelo RH.')] if TESTE or not str(x[0]).startswith('DEMONSTRACAO')]:
    line(a,b)
ws.sheet_view.showGridLines=False

# ---------- PAINEL (visual moderno) ----------
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import DataBarRule
from openpyxl.worksheet.properties import PageSetupProperties
ws_pn=wb.create_sheet('PAINEL'); ws_pn.sheet_properties.tabColor='0F172A'
ws=ws_pn
ws.sheet_view.showGridLines=False
NAVY='0F172A'; BG='F1F5F9'; WHITE='FFFFFF'; SLATE='334155'; MUTED='64748B'; LINE='E2E8F0'
BLUE='2563EB'; INDIGO='4F46E5'; VIOLET='7C3AED'; AMBER='D97706'; RED='DC2626'; GREEN='16A34A'; GRAY='475569'
FN='Calibri'
def F(size=10,bold=False,color=SLATE,italic=False): return Font(name=FN,size=size,bold=bold,color=color,italic=italic)
def fill(c): return PatternFill('solid',fgColor=c)
LASTCOL=17   # Q
ws.column_dimensions['A'].width=2
for c in range(2,LASTCOL+1): ws.column_dimensions[L(c)].width=12.6
ws.column_dimensions['R'].width=2
# fundo cinza claro em toda a área
for r in range(1,140):
    for c in range(1,LASTCOL+2): ws.cell(r,c).fill=fill(BG)
RT=RE_LAST+1
RM=lambda col: f'RESUMO_MENSAL!${col}${RT}'
PENDCNT=f'COUNT({cr("AK")})'
def paint(r1,c1,r2,c2,color):
    for r_ in range(r1,r2+1):
        for c_ in range(c1,c2+1): ws.cell(r_,c_).fill=fill(color)
def merge(r1,c1,r2,c2):
    if (r1,c1)!=(r2,c2): ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)

# ---- banner
paint(1,2,3,LASTCOL,NAVY)
ws.row_dimensions[1].height=30; ws.row_dimensions[2].height=18; ws.row_dimensions[3].height=8
merge(1,2,1,11); ws['B1']='Painel Mensal do Ponto'; ws['B1'].font=F(22,True,WHITE); ws['B1'].alignment=Alignment(vertical='center')
merge(2,2,2,11)
ws['B2']='="'+EMPNOME+'  •  Competência "&RIGHT(P_COMP,2)&"/"&LEFT(P_COMP,4)&"  •  Apuração "&TEXT(P_DT_INI,"dd/mm")&" a "&TEXT(P_DT_FIM,"dd/mm/yyyy")&"  •  Jornada "&TEXT(P_JORNADA,"hh:mm")&" (seg-sex)"'
ws['B2'].font=F(10,False,'94A3B8')
merge(1,13,2,LASTCOL); ws['M1']='=EXPORT_CALC!V11'
ws['M1'].font=F(11,True,WHITE); ws['M1'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for c in range(13,LASTCOL+1):
    for r_ in (1,2): ws.cell(r_,c).fill=fill('475569')
ws.conditional_formatting.add('M1:Q2',FormulaRule(formula=['LEFT($M$1,3)="NÃO"'],fill=PatternFill('solid',bgColor=RED,fgColor=RED)))
ws.conditional_formatting.add('M1:Q2',FormulaRule(formula=['LEFT($M$1,3)="PRO"'],fill=PatternFill('solid',bgColor=GREEN,fgColor=GREEN)))
ws.row_dimensions[4].height=8
# ---- faixa de atenção (preenchida ao final)
merge(5,2,5,LASTCOL); ws.row_dimensions[5].height=26
ws.row_dimensions[6].height=8

# ---- cartões de indicadores
def card(r,c,label,formula,caption,color,fmt=None):
    for rr_ in (r,r+1,r+2):
        merge(rr_,c,rr_,c+3)
        for cc_ in range(c,c+4):
            cell=ws.cell(rr_,cc_); cell.fill=fill(WHITE)
            cell.border=Border(left=Side('thick',color=BG) if cc_==c else None,right=Side('thick',color=BG) if cc_==c+3 else None,
                               top=Side('thick',color=color) if rr_==r else None,bottom=Side('thick',color=BG) if rr_==r+2 else None)
    a=ws.cell(r,c,label); a.font=F(9,True,MUTED); a.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    b_=ws.cell(r+1,c,formula); b_.font=F(26,True,color); b_.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    if fmt: b_.number_format=fmt
    d=ws.cell(r+2,c,caption); d.font=F(9,False,'94A3B8'); d.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    ws.row_dimensions[r].height=22; ws.row_dimensions[r+1].height=40; ws.row_dimensions[r+2].height=20
NPONTO='SUMPRODUCT(--(LEFT(RESUMO_STATUS,8)<>"IGNORADO"))'
card(7,2,'FUNCIONÁRIOS NO PONTO','='+NPONTO,'com controle de ponto neste mês',BLUE,'0')
card(7,6,'HORAS EXTRAS (70% + 100%)',f'={RM("S")}+{RM("T")}',f'="a 70%: "&TEXT({RM("S")},"[h]:mm")&"   •   a 100%: "&TEXT({RM("T")},"[h]:mm")',INDIGO,T_HMZ)
card(7,10,'ADICIONAL NOTURNO',f'={RM("N")}','horas reais entre 22h e 5h',VIOLET,T_HMZ)
card(7,14,'FALTAS / ATRASOS (HORAS)',f'={RM("Q")}',f'="+ "&{RM("E")}&" falta(s) de dia inteiro"',GRAY,T_HMZ)
card(11,2,'A CONFIRMAR COM A EMPRESA','='+PENDCNT,'itens em aberto (lista abaixo)',AMBER,'0')
card(11,6,'DIAS COM IRREGULARIDADE LEGAL',f'=SUM({cr("AS")})','ver seção A',RED,'0')
card(11,10,'DSR DESCONTADO (SEMANAS)',f'={RM("P")}','falta de meio período ou mais',GRAY,'0')
card(11,14,'SEM MARCAÇÃO E SEM OBSERVAÇÃO',f'=COUNT({cr("AT")})','dias úteis: perguntar ao cliente',AMBER,'0')
ws.row_dimensions[10].height=4; ws.row_dimensions[14].height=8

# ---- utilitários de seção / tabela
def section(r,text,color,badge=None,bw=3):
    merge(r,2,r,LASTCOL-bw)
    ws.cell(r,2,text).font=F(13,True,NAVY); ws.cell(r,2).alignment=Alignment(vertical='center')
    for c in range(2,LASTCOL+1): ws.cell(r,c).border=Border(bottom=Side('medium',color=color))
    if badge:
        merge(r,LASTCOL-bw+1,r,LASTCOL); b_=ws.cell(r,LASTCOL-bw+1,badge); b_.font=F(11,True,color); b_.alignment=Alignment(horizontal='right',vertical='center')
    ws.row_dimensions[r].height=26
def th(r,items):
    for (c1,c2,t) in items:
        merge(r,c1,r,c2); c=ws.cell(r,c1,t); c.font=F(9,True,MUTED); c.alignment=Alignment(horizontal='left' if c1==items[0][0] else 'center',vertical='center',wrap_text=True,indent=1 if c1==items[0][0] else 0)
        for cc in range(c1,c2+1): ws.cell(r,cc).fill=fill('F8FAFC'); ws.cell(r,cc).border=Border(bottom=Side('thin',color='CBD5E1'))
    ws.row_dimensions[r].height=24
def td(r,items,fmts=None,zebra=False,h=24):
    paint(r,2,r,LASTCOL,'F8FAFC' if zebra else WHITE)
    for i,(c1,c2,v) in enumerate(items):
        merge(r,c1,r,c2); c=ws.cell(r,c1,v); c.font=F(10); c.alignment=Alignment(horizontal='left' if i==0 else 'center' if (c2==c1 or (c2-c1)<=1) else 'left',vertical='center',wrap_text=True,indent=1 if i==0 or (c2-c1)>1 else 0)
        if fmts and fmts.get(c1): c.number_format=fmts[c1]
    for cc in range(2,LASTCOL+1): ws.cell(r,cc).border=Border(bottom=Side('thin',color=LINE))
    ws.row_dimensions[r].height=h
def cf_pill(rng,key_formula,kind):
    pal={'red':('FEE2E2','B91C1C'),'amber':('FEF3C7','B45309'),'green':('DCFCE7','15803D')}[kind]
    ws.conditional_formatting.add(rng,FormulaRule(formula=[key_formula],font=Font(bold=True,color=pal[1]),fill=PatternFill('solid',bgColor=pal[0],fgColor=pal[0])))

# ---- Pendências para confirmar com o cliente (lista única, com links)
_est=0
for _x in recs:
    if DT_INI<=_x['date']<=DT_FIM:
        _n=len(_x['punches']); _nt=(_x['note'] or '').upper()
        if _n%2==1: _est+=1
        elif _n==0 and _x['date'].weekday()<5 and not _nt: _est+=1
        elif _n>0 and any(_t in _nt for _t in ('FÉRIAS','AFASTAMENTO','LICENÇA')): _est+=1
PN=min(120,max(15,_est+10))
r=16
section(r,'Pendências para confirmar com o cliente',RED,f'=COUNT({cr("AK")})&" item(ns) em aberto"',bw=4)
r+=1; merge(r,2,r,LASTCOL)
ws.cell(r,2,'Clique no nome para ir direto à célula do PONTO_BRUTO que precisa ser corrigida. Depois de corrigir (e registrar o motivo na coluna K), o item some desta lista.').font=F(9,False,MUTED,True)
ws.cell(r,2).alignment=Alignment(vertical='center'); ws.row_dimensions[r].height=20
r+=1; merge(r,2,r,LASTCOL)
_tp=lambda t: f'COUNTIF(PENDENCIAS!$F$6:$F$125,"{t}")'
ws.cell(r,2,f'="Marcação ímpar: "&{_tp("Marcação ímpar")}&"     •     Dia útil sem marcação e sem observação: "&{_tp("Dia útil sem marcação")}&"     •     Marcação em férias/afastamento: "&{_tp("Marcação em férias/afastamento")}&"     •     Ocorrência não cadastrada: "&{_tp("Ocorrência não cadastrada")}').font=F(10,True,NAVY)
ws.cell(r,2).alignment=Alignment(vertical='center'); ws.row_dimensions[r].height=22
r+=1; th(r,[(2,5,'FUNCIONÁRIO (clique)'),(6,7,'DATA'),(8,9,'MARCAÇÕES'),(10,12,'TIPO'),(13,LASTCOL,'O QUE CONFIRMAR')])
PD0=r+1
for k in range(1,PN+1):
    r+=1; pr=5+k; idx=f'PENDENCIAS!$I${pr}'
    colx=f'IF(PENDENCIAS!$F${pr}="Marcação ímpar",CHOOSE((INDEX({cr("J")},{idx})+1)/2,"E","G","I"),"J")'
    name_f=f'=IF({idx}="","",HYPERLINK("#PONTO_BRUTO!"&{colx}&({FIRST-1}+{idx}),PENDENCIAS!$B${pr}))'
    td(r,[(2,5,name_f),(6,7,f'=IF({idx}="","",TEXT(PENDENCIAS!$C${pr},"dd/mm")&" ("&PENDENCIAS!$D${pr}&")")'),(8,9,f'=IF({idx}="","",IF(PENDENCIAS!$E${pr}="","sem marcações",PENDENCIAS!$E${pr}))'),(10,12,f'=IF({idx}="","",PENDENCIAS!$F${pr})'),(13,LASTCOL,f'=IF({idx}="","",PENDENCIAS!$G${pr})')],zebra=(k%2==0),h=28)
    ws.cell(r,2).font=Font(name=FN,size=10,bold=True,color=BLUE,underline='single')
PD1=r
ws.conditional_formatting.add(f'B{PD0}:Q{PD1}',FormulaRule(formula=[f'$B{PD0}=""'],fill=PatternFill('solid',bgColor=BG,fgColor=BG),border=Border(bottom=Side(style='thin',color=BG))))
cf_pill(f'J{PD0}:L{PD1}',f'$J{PD0}="Dia útil sem marcação"','red')
cf_pill(f'J{PD0}:L{PD1}',f'AND($J{PD0}<>"",$J{PD0}<>"Dia útil sem marcação")','amber')
r+=1; merge(r,2,r,LASTCOL)
ws.cell(r,2,f'=IF(COUNT({cr("AK")})>{PN},"+ "&(COUNT({cr("AK")})-{PN})&" pendência(s) além das "&{PN}&" exibidas: corrija as primeiras e a lista se atualiza.",IF(COUNT({cr("AK")})=0,"Nenhuma pendência: nada a perguntar ao cliente neste mês.","Se o cliente confirmar falta, registre FALTA INJUSTIFICADA na coluna J do PONTO_BRUTO; folga, férias ou atestado também se registram ali."))').font=F(9,False,MUTED,True)
ws.row_dimensions[r].height=22
paint(r+1,2,r+1,LASTCOL,BG); ws.row_dimensions[r+1].height=10; r+=1

# ---- B. fora da lei
r+=1; section(r,'A  Fora da lei: risco trabalhista',RED,f'="Dias irregulares: "&SUM({cr("AS")})')
r+=1; th(r,[(2,6,'IRREGULARIDADE'),(7,9,'BASE LEGAL'),(10,10,'DIAS'),(11,11,'FUNC.'),(12,12,'SITUAÇÃO'),(13,LASTCOL,'OBSERVAÇÃO')])
irr=[('Jornada acima de 10h no dia','CLT art.59','AM','Excesso de horas extras diárias (máx. 2h/dia).'),
     ('Interjornada menor que 11h','CLT art.66','AN','Descanso insuficiente entre jornadas.'),
     ('Almoço não registrado (jornada > 6h)','CLT art.71','AO','Pago como hora extra (regra da empresa); risco de autuação.'),
     ('Intervalo registrado menor que 1h','CLT art.71 / CCT','AP','Só vale se norma coletiva/MTE permitir redução.'),
     ('Intervalo maior que 2h','CLT art.71','AQ','Exige acordo escrito/norma coletiva.'),
     ('Trabalho em férias/afastamento','CLT (férias)','AR','Consultar o jurídico sobre pagamento em dobro.')]
B0=r+1; Brows={}
for k,(lab,base,col,obs) in enumerate(irr):
    r+=1; Brows[col]=r
    td(r,[(2,6,lab),(7,9,base),(10,10,f'=SUM({cr(col)})'),(11,11,None),(12,12,f'=IF(J{r}=0,"OK","IRREGULAR")'),(13,LASTCOL,obs)],{10:'0',11:'0'},zebra=(k%2==1),h=26)
B1=r
r+=1; BT=r
td(r,[(2,9,'Total de dias com alguma irregularidade (um dia pode ter mais de uma)'),(10,10,f'=SUM({cr("AS")})'),(11,11,None),(12,12,f'=IF(J{r}=0,"OK","IRREGULAR")'),(13,LASTCOL,'Detalhe dia a dia na aba oculta CALCULO_DIARIO (coluna Alertas; clique direito em uma aba > Reexibir).')],{10:'0',11:'0'},h=28)
ws.cell(r,2).font=F(10,True,NAVY); ws.cell(r,10).font=F(11,True,RED)
cf_pill(f'L{B0}:L{BT}',f'L{B0}="IRREGULAR"','red'); cf_pill(f'L{B0}:L{BT}',f'L{B0}="OK"','green')
ws.conditional_formatting.add(f'J{B0}:J{B1}',DataBarRule(start_type='num',start_value=0,end_type='max',color='FCA5A5'))
paint(r+1,2,r+1,LASTCOL,BG); ws.row_dimensions[r+1].height=10; r+=1

# ---- C. ocorrências
r+=1; section(r,'B  Ocorrências do mês',BLUE)
r+=1; th(r,[(2,7,'OCORRÊNCIA'),(8,9,'DIAS ÚTEIS'),(10,LASTCOL,'COMO É TRATADA')])
occ=[('Atestado (inclui parcial)','ATESTADO','Abona a falta; sem desconto de salário nem DSR (CCT cl.35).'),
     ('Férias','FÉRIAS','Sem falta. Lançar no módulo Férias do Domínio.'),
     ('Licença','LICENÇA','Sem desconto (ex.: paternidade).'),
     ('Falta justificada','FALTA JUST.','CCT cl.28.'),
     ('Afastamento (INSS/acidente)','AFASTAMENTO','Lançar no Domínio.'),
     ('Compensação / abono','COMPENSAÇÃO','Sem desconto.')]
for k,(lab,key,tx) in enumerate(occ):
    r+=1; td(r,[(2,7,lab),(8,9,f'=COUNTIFS({cr("F")},"{key}",{cr("I")},">0")'),(10,LASTCOL,tx)],{8:'0'},zebra=(k%2==1))
r+=1; td(r,[(2,7,'Faltas de dia inteiro sem ocorrência (desconto + perda de DSR)'),(8,9,f'=SUM({cr("AA")})'),(10,LASTCOL,'Falta injustificada: desconta o dia (07:20 no Domínio) e o DSR da semana.')],{8:'0'})
ws.cell(r,2).font=F(10,True,NAVY); cf_pill(f'H{r}',f'H{r}>0','red')
r+=1; td(r,[(2,7,'DSR descontado (semanas com falta de meio período ou mais)'),(8,9,f'=RESUMO_MENSAL!$P${RT}'),(10,LASTCOL,'Regra da empresa: meio período de falta injustificada no dia desconta o DSR da semana.')],{8:'0'},zebra=True)
ws.cell(r,2).font=F(10,True,NAVY); cf_pill(f'H{r}',f'H{r}>0','red')
paint(r+1,2,r+1,LASTCOL,BG); ws.row_dimensions[r+1].height=10; r+=1

# ---- D. qualidade do software
r+=1; section(r,'C  Qualidade do registro e do software do relógio',GREEN)
r+=1; th(r,[(2,8,'INDICADOR'),(9,10,'DIAS'),(11,LASTCOL,'POR QUE IMPORTA')])
pbr=lambda col: f'PONTO_BRUTO!${col}${FIRST}:${col}${LAST}'
qual=[('Software com jornada diferente da real (CHPrev x jornada)',f'=SUMPRODUCT(({pbr("L")}>0)*(ABS({pbr("L")}-P_JORNADA)>0.0001)*({cr("H")}=1))','Relógio configurado com outra jornada: cálculos do software saem errados.'),
      ('Falta de dia inteiro do software com marcações completas',f'=SUMPRODUCT(({pbr("L")}>0)*({pbr("N")}>={pbr("L")}-0.0001)*({cr("K")}="OK"))','Descontar isso na folha seria cobrança indevida.'),
      ('Falta do software em dia com marcação ímpar',f'=SUMPRODUCT(({pbr("N")}>0)*({cr("K")}="MARCAÇÃO ÍMPAR"))','Provável esquecimento de bater; não é falta comprovada.'),
      ('Adicional noturno não apurado pelo software',f'=SUMPRODUCT(({cr("X")}>0)*({pbr("S")}=0))','Direito do trabalhador não calculado pelo relógio.'),
      ('Dias com só 1 par de marcações (sem almoço batido)',f'=SUM({cr("AO")})','Falta de controle do intervalo.')]
Q0=r+1
for k,(lab,fx,why) in enumerate(qual):
    r+=1; td(r,[(2,8,lab),(9,10,fx),(11,LASTCOL,why)],{9:'0'},zebra=(k%2==1),h=26)
cf_pill(f'I{Q0}:I{r}',f'I{Q0}>0','amber'); cf_pill(f'I{Q0}:I{r}',f'I{Q0}=0','green')
paint(r+1,2,r+1,LASTCOL,BG); ws.row_dimensions[r+1].height=10; r+=1

# ---- E. por funcionário
r+=1; section(r,'D  Resultado por funcionário',NAVY,'vermelho = irregularidade  •  amarelo = pendência',bw=6)
r+=1
heads=['STATUS','HE 70%','HE 100%','FALTAS PARCIAIS','FALTAS INTEIRAS','NOTURNO','A CONFIRMAR','JORNADA > 10H','INTERJ. < 11H','SEM ALMOÇO','INTERV. < 1H / > 2H','TRAB. EM FÉRIAS','TOTAL IRREG.','DSR DESCONT.']
ws.row_dimensions[r].height=34
merge(r,2,r,3); c=ws.cell(r,2,'FUNCIONÁRIO'); c.font=F(9,True,MUTED); c.alignment=Alignment(vertical='center',indent=1)
for cc in range(2,LASTCOL+1): ws.cell(r,cc).fill=fill('F8FAFC'); ws.cell(r,cc).border=Border(bottom=Side('thin',color='CBD5E1'))
for i,h_ in enumerate(heads):
    c=ws.cell(r,4+i,h_); c.font=F(8,True,MUTED); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
E0=r+1
for i,e in enumerate(emps):
    r+=1; rr=RE_FIRST+i
    paint(r,2,r,LASTCOL,'F8FAFC' if i%2 else WHITE); ws.row_dimensions[r].height=22
    merge(r,2,r,3); ws.cell(r,2,f'=IFERROR(HYPERLINK("#PONTO_BRUTO!D"&({FIRST-1}+MATCH(RESUMO_MENSAL!$A${rr},PONTO_BRUTO!$A${FIRST}:$A${LAST},0)),RESUMO_MENSAL!$A${rr}),RESUMO_MENSAL!$A${rr})').font=Font(name=FN,size=10,bold=True,color=BLUE,underline='single'); ws.cell(r,2).alignment=Alignment(vertical='center',indent=1)
    A_=f'{cr("A")},$B{r}'
    Rr=lambda col: f'=RESUMO_MENSAL!${col}${rr}'
    vals=[('D',f'=IF(LEFT(RESUMO_MENSAL!$Y${rr},2)="OK","OK",LEFT(RESUMO_MENSAL!$Y${rr},FIND(" ",RESUMO_MENSAL!$Y${rr}&" ")-1))',None),
          ('E',Rr('S'),T_HM),('F',Rr('T'),T_HM),('G',Rr('Q'),T_HM),('H',Rr('E'),'0'),('I',Rr('N'),T_HM),
          ('J',f'=SUMIFS({cr("AJ")},{A_})','0'),('K',f'=SUMIFS({cr("AM")},{A_})','0'),('L',f'=SUMIFS({cr("AN")},{A_})','0'),('M',f'=SUMIFS({cr("AO")},{A_})','0'),
          ('N',f'=SUMIFS({cr("AP")},{A_})+SUMIFS({cr("AQ")},{A_})','0'),('O',f'=SUMIFS({cr("AR")},{A_})','0'),
          ('P',f'=SUMIFS({cr("AS")},{A_})','0'),('Q',f'=RESUMO_MENSAL!$P${rr}','0')]
    for (col,fx,fm) in vals:
        c=ws[f'{col}{r}']; c.value=fx; c.font=F(10); c.alignment=Alignment(horizontal='center',vertical='center')
        if fm: c.number_format=fm
    for cc in range(2,LASTCOL+1): ws.cell(r,cc).border=Border(bottom=Side('thin',color=LINE))
E1r=r
paint(r+1,2,r+1,LASTCOL,BG); ws.row_dimensions[r+1].height=10
for col,brow in Brows.items():
    ws.cell(brow,11).value=f'=SUMPRODUCT(--(COUNTIFS({cr("A")},$B${E0}:$B${E1r},{cr(col)},1)>0))'
ws.cell(BT,11).value=f'=COUNTIF(P{E0}:P{E1r},">0")'
cf_pill(f'D{E0}:D{E1r}',f'D{E0}="BLOQUEADO"','red'); cf_pill(f'D{E0}:D{E1r}',f'D{E0}="REVISAR"','amber'); cf_pill(f'D{E0}:D{E1r}',f'D{E0}="OK"','green'); cf_pill(f'D{E0}:D{E1r}',f'D{E0}="IGNORADO"','amber')
ws.conditional_formatting.add(f'K{E0}:P{E1r}',CellIsRule(operator='greaterThan',formula=['0'],font=Font(bold=True,color='B91C1C'),fill=PatternFill('solid',bgColor='FEE2E2',fgColor='FEE2E2')))
ws.conditional_formatting.add(f'J{E0}:J{E1r}',CellIsRule(operator='greaterThan',formula=['0'],font=Font(bold=True,color='B45309'),fill=PatternFill('solid',bgColor='FEF3C7',fgColor='FEF3C7')))
ws.conditional_formatting.add(f'E{E0}:F{E1r}',DataBarRule(start_type='num',start_value=0,end_type='max',color='93C5FD'))

# ---- faixa de atenção
ws['B5']=f'=IF(AND({PENDCNT}=0,SUM({cr("AS")})=0),"Tudo em dia: nenhuma pendência e nenhuma irregularidade neste mês.","Atenção: "&{PENDCNT}&" item(ns) a confirmar com a empresa"&IF(COUNT({cr("AT")})>0," (inclui "&COUNT({cr("AT")})&" dia(s) útil(eis) sem marcação e sem observação)","")&" e "&SUM({cr("AS")})&" dia(s) com irregularidade legal em "&K{BT}&" funcionário(s).")'
ws['B5'].font=F(11,True,'92400E'); ws['B5'].alignment=Alignment(vertical='center',indent=1)
paint(5,2,5,LASTCOL,'FEF3C7')
ws.conditional_formatting.add('B5:Q5',FormulaRule(formula=[f'AND({PENDCNT}=0,SUM({cr("AS")})=0)'],font=Font(bold=True,color='15803D'),fill=PatternFill('solid',bgColor='DCFCE7',fgColor='DCFCE7')))

# ---- F. gráficos + dados auxiliares (colunas T:AC)
r+=1; section(r,'E  Gráficos',NAVY)
CH0=r+1
H0=20  # coluna T
ws.column_dimensions['S'].width=2
ws.cell(3,H0,'Dados auxiliares dos gráficos (não editar)').font=F(8,False,'94A3B8',True)
for i,h_ in enumerate(['Funcionário','HE 70% (h)','HE 100% (h)','Chave','Posição','','Ranking','Funcionário','HE 70% (h)','HE 100% (h)']):
    if h_: ws.cell(4,H0+i,h_).font=F(8,True,'94A3B8')
cN,c70,c100,cK,cP,cR,cN2,c702,c1002=H0,H0+1,H0+2,H0+3,H0+4,H0+6,H0+7,H0+8,H0+9
NEMP=NE
for i in range(NEMP):
    rr=RE_FIRST+i; r_=5+i
    ws.cell(r_,cN,f'=RESUMO_MENSAL!$A${rr}')
    ws.cell(r_,c70,f'=RESUMO_MENSAL!$S${rr}*24').number_format='0.00'
    ws.cell(r_,c100,f'=RESUMO_MENSAL!$T${rr}*24').number_format='0.00'
    ws.cell(r_,cK,f'={L(c70)}{r_}+{L(c100)}{r_}+ROW()/1000000')
    ws.cell(r_,cP,f'=RANK({L(cK)}{r_},${L(cK)}$5:${L(cK)}${4+NEMP},0)')
    ws.cell(r_,cR,i+1)
    ws.cell(r_,cN2,f'=INDEX(${L(cN)}$5:${L(cN)}${4+NEMP},MATCH({L(cR)}{r_},${L(cP)}$5:${L(cP)}${4+NEMP},0))')
    ws.cell(r_,c702,f'=INDEX(${L(c70)}$5:${L(c70)}${4+NEMP},MATCH({L(cR)}{r_},${L(cP)}$5:${L(cP)}${4+NEMP},0))').number_format='0.00'
    ws.cell(r_,c1002,f'=INDEX(${L(c100)}$5:${L(c100)}${4+NEMP},MATCH({L(cR)}{r_},${L(cP)}$5:${L(cP)}${4+NEMP},0))').number_format='0.00'
    for cc in range(H0,H0+10): ws.cell(r_,cc).font=F(8,False,'94A3B8')
ws.cell(4,cN2,'Funcionário'); ws.cell(4,c702,'HE 70% (h)'); ws.cell(4,c1002,'HE 100% (h)')
def style_chart(ch,colors):
    ch.graphical_properties=GraphicalProperties(ln=LineProperties(noFill=True))
    for s,col in zip(ch.series,colors):
        s.graphicalProperties.solidFill=col; s.graphicalProperties.line.solidFill=col
    ch.gapWidth=55
    ch.x_axis.delete=False; ch.y_axis.delete=False; ch.x_axis.tickLblSkip=1
    ch.y_axis.majorGridlines=None
    ch.x_axis.scaling.orientation='maxMin'
ch1=BarChart(); ch1.type='bar'; ch1.grouping='stacked'; ch1.overlap=100
ch1.title='Horas extras por funcionário (h)'; ch1.height=11.5; ch1.width=19
ch1.add_data(Reference(ws,min_col=c702,max_col=c1002,min_row=4,max_row=4+NEMP),titles_from_data=True); ch1.set_categories(Reference(ws,min_col=cN2,min_row=5,max_row=4+NEMP))
ch1.y_axis.scaling.min=0; ch1.legend.position='b'
style_chart(ch1,[BLUE,'F59E0B'])
ws.add_chart(ch1,f'B{CH0}')
ch2=BarChart(); ch2.type='bar'; ch2.title='Irregularidades legais (dias)'; ch2.height=11.5; ch2.width=19
ch2.add_data(Reference(ws,min_col=10,min_row=B0-1,max_row=B1),titles_from_data=True); ch2.set_categories(Reference(ws,min_col=2,min_row=B0,max_row=B1)); ch2.legend=None
style_chart(ch2,[RED])
ws.add_chart(ch2,f'J{CH0}')
paint(CH0,2,CH0+23,LASTCOL,BG)
ws.freeze_panes='A4'
ws.print_area=f'A1:R{CH0+24}'
ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0

wb.move_sheet('LEIA_ME',offset=-wb.index(wb['LEIA_ME']))
wb.move_sheet('PAINEL',offset=-wb.sheetnames.index('PAINEL'))
for _s in ('CALCULO_DIARIO','SEMANAL','EXPORT_CALC','PENDENCIAS'): wb[_s].sheet_state='hidden'
wb.active=0
if not TESTE and 'DEMONSTRACAO' in wb.sheetnames: wb.remove(wb['DEMONSTRACAO'])
wb.save(CFG['saida'])
print(f'Planilha gerada: {CFG["saida"]}')
print(f'  {NE} funcionário(s), {len(recs)} dia(s), apuração {DT_INI:%d/%m/%Y} a {DT_FIM:%d/%m/%Y}, competência {COMP[4:]}/{COMP[:4]}')
print('  Feriados no período: '+(', '.join(f'{d:%d/%m} {n}' for d,n in FERIADOS.items()) or 'nenhum'))
for a in AVISOS: print('  AVISO: '+a)
